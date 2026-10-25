"""Write an inlined document as a Word file, from bytes only.

The input is HTML the inliner already cleaned: every image it keeps is a
``data:`` URI. The writer parses it with network access off, walks the tree
with an explicit stack, and builds the package in memory with python-docx.
python-docx opens any path string it is handed, so images reach it only as
bytes, and only PNG, JPEG and GIF reach it at all. Nothing here decodes
pixels.
"""
from __future__ import annotations

import base64
import binascii
import copy
import io
import itertools
import re
import time
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime
from urllib.parse import quote, urlsplit

import docx
import lxml.html
from docx.document import Document as WordDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.image.image import Image
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml.numbering import CT_Num
from docx.oxml.parser import OxmlElement
from docx.oxml.simpletypes import ST_Merge
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Emu, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml.etree import ParserError
from lxml.html import HtmlElement

from cowork.schemas.artifact_export import InlinedDocument
from cowork.services.artifact_inline import (
    GIF_SIGNATURES,
    JPEG_SIGNATURE,
    JPEG_STANDALONE_MARKERS,
    PNG_SIGNATURE,
    TRUNCATION_NOTE,
    XML_UNSAFE,
    UriPlaceholders,
    parse_stopped,
)

MAX_DEPTH = 256
MAX_NODES = 200_000
DEADLINE_SECONDS = 10.0
# Word's own column limit, and a cap on every table cell in one document.
MAX_TABLE_COLUMNS = 63
MAX_TABLE_CELLS = 20_000
# Each nested table adds three levels to document.xml (table, row and cell),
# and Word's and python-docx's XML parsers refuse a part nested past 256
# levels. A table nested deeper than this lets its cells' text flow into the
# cell around it.
MAX_TABLE_NESTING = 32
MAX_LIST_LEVEL = 3
# Word's highest list start number.
MAX_LIST_NUMBER = 32_767
# python-docx refuses a core property longer than 255 characters.
MAX_TITLE_LENGTH = 255
LINK_SCHEMES = frozenset({"http", "https", "mailto"})
MONOSPACE_FONT = "Courier New"
# The paragraph that ends a document whose tables lost cells to the limits above.
TABLE_CELLS_NOTE = "Some table cells were left out because a table was too large to convert."

_LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
_SMALL_SIZE = Pt(9)
_EMU_PER_INCH = 914_400
_DEFAULT_DPI = 72

_HTML_SPACE = re.compile("[ \t\n\r\f]+")
# What RFC 3986 lets a URI hold as it is, beyond letters, digits and "-._~".
# Every other character in a link target is percent-encoded, so a consumer
# that checks a relationship's Target reads a valid URI.
_URI_SAFE = "!#$%&'()*+,/:;=?@[]"
# A "%" that starts no escape, which the encoding writes as "%25".
_LONE_PERCENT = re.compile("%(?![0-9A-Fa-f]{2})")
_LEADING_DIGITS = re.compile(r"\s*([0-9]+)")
# python-docx writes each of these as its own element.
_PRE_BREAKS = re.compile("[\n\t]")
# A character the parser keeps as it is and that can't end a URI: no
# whitespace, quote, bracket or entity. A data: URI of 1024 or more of them
# is set aside before parsing; matching only at the first character of a
# run keeps the scan linear.
_URI_CHARACTER = r"[^\s\"'<>&()]"
_LONG_DATA_URI = re.compile(rf"(?<!{_URI_CHARACTER})(?i:data:){_URI_CHARACTER}{{1024,}}")
# The attributes the writer reads text or a URI from.
_READ_ATTRIBUTES = ("src", "alt", "href")

# python-docx keeps every chunk of a PNG, every marker before a JPEG's scan,
# and every entry of each Exif block's first directory, strings and all,
# before it reads an image's size. So a small crafted header can cost a
# second and hundreds of MiB, twice per image. No camera or chart library
# writes a header past these; an image past them becomes its placeholder
# text. The marker cap counts each 0xFF 0x00 pair among the markers too, as
# python-docx steps past those one at a time.
_MAX_PNG_CHUNKS = 4096
_MAX_JPEG_MARKERS = 4096
_MAX_EXIF_ENTRIES = 256
_JPEG_EOI, _JPEG_SOS, _JPEG_APP1 = 0xD9, 0xDA, 0xE1

# Elements whose content is not document text.
_SKIPPED_TAGS = frozenset(
    {
        "head", "title", "script", "style", "template", "svg", "math",
        "iframe", "object", "embed", "canvas", "audio", "video",
    }
)
_HEADING_LEVELS = {f"h{level}": level for level in range(1, 7)}
# Blocks that start their own paragraph as soon as they open.
_PARAGRAPH_TAGS = frozenset(
    {"p", "pre", "li", "dt", "dd", "figcaption", "summary", "legend", *_HEADING_LEVELS}
)
# Blocks that only end the paragraph before and after them.
_CONTAINER_TAGS = frozenset(
    {
        "html", "body", "div", "section", "article", "header", "footer", "main",
        "nav", "aside", "blockquote", "figure", "form", "fieldset", "details",
        "dl", "hr", "address", "center", "ul", "ol", "table", "tr", "td", "th",
        "thead", "tbody", "tfoot", "caption",
    }
)
# A list item's first block reuses the item's empty paragraph, so a loose
# list (each item wrapped in <p>) keeps its bullet on the text.
_ITEM_REUSING_TAGS = frozenset({"p", "div", "section", "article"})
_TABLE_ROW_GROUPS = frozenset({"thead", "tbody", "tfoot"})
_TABLE_CELLS = frozenset({"td", "th"})
# Elements a browser keeps in a table, even inside a wrapper such as form.
_TABLE_STRUCTURE = ("tr", "td", "th", "thead", "tbody", "tfoot")
# The <ol type> values Word can show, as Word number formats.
_NUMBER_FORMATS = {"a": "lowerLetter", "A": "upperLetter", "i": "lowerRoman", "I": "upperRoman"}


@dataclass(frozen=True, eq=False)
class _Link:
    """One <a> element's allowed target; runs compare links by identity."""

    url: str


@dataclass(frozen=True)
class _RunStyle:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    monospace: bool = False
    subscript: bool = False
    superscript: bool = False
    highlight: bool = False
    small: bool = False
    link: _Link | None = None


_INLINE_STYLES: dict[str, dict[str, bool]] = {
    "b": {"bold": True},
    "strong": {"bold": True},
    "i": {"italic": True},
    "em": {"italic": True},
    "u": {"underline": True},
    "ins": {"underline": True},
    "s": {"strike": True},
    "strike": {"strike": True},
    "del": {"strike": True},
    "code": {"monospace": True},
    "kbd": {"monospace": True},
    "samp": {"monospace": True},
    "tt": {"monospace": True},
    "sub": {"subscript": True},
    "sup": {"superscript": True},
    "mark": {"highlight": True},
    "small": {"small": True},
}


@dataclass(frozen=True)
class _OpenLink:
    link: _Link
    element: BaseOxmlElement


@dataclass
class _Flow:
    """The paragraph that text lands in, for the body or for one table cell."""

    container: WordDocument | _Cell
    width: int
    paragraph: Paragraph | None = None
    # A table cell starts with one empty paragraph; the cell's first text uses it.
    reusable: Paragraph | None = None
    open_link: _OpenLink | None = None
    at_space: bool = True
    item_pending: bool = False


@dataclass(eq=False)
class _ListNumbering:
    """One list's level, style and numbering as its items go by.

    Items share one w:num and Word counts them on. An item's value attribute
    starts a new w:num at that number, and in a reversed list every item
    gets its own w:num, since Word only counts up.
    """

    level: int  # 1 to MAX_LIST_LEVEL
    style: str
    ordered: bool
    number_format: str | None
    reversed: bool
    next_number: int
    num_id: int | None = None


@dataclass(frozen=True)
class _Context:
    flow: _Flow
    run: _RunStyle = field(default_factory=_RunStyle)
    pre: bool = False
    paragraph_style: str | None = None
    list_numbering: _ListNumbering | None = None  # the innermost list around the element
    depth: int = 0
    table_depth: int = 0


@dataclass(frozen=True)
class _TableParts:
    """A table's rows of cells, its captions, and the content a browser shows before it."""

    rows: list[list[HtmlElement]]
    captions: list[HtmlElement]
    strays: list[HtmlElement]


@dataclass(frozen=True)
class _TableItem:
    """One element met while sorting a table's content, and the row it sits in.

    Like a walk visit, an item with ``siblings`` set queues the element's next
    sibling when it is taken.
    """

    element: HtmlElement
    row: list[HtmlElement] | None
    leaving: bool = False
    siblings: bool = False


@dataclass(frozen=True)
class _Visit:
    element: HtmlElement
    outer: _Context
    leaving: bool = False
    write_tail: bool = True
    prepaid: bool = False
    # Set on the visit that adds a table once the content before it is written.
    table_parts: _TableParts | None = None
    # Queue the element's next sibling when this visit is taken, so the stack
    # holds one pending visit per open element instead of every child at once.
    siblings: bool = False


@dataclass(frozen=True)
class _PictureSize:
    width: int
    height: int


@dataclass(frozen=True)
class _CellPlacement:
    element: HtmlElement
    row: int
    column: int
    rowspan: int
    colspan: int


@dataclass(frozen=True)
class _TableLayout:
    placements: list[_CellPlacement]
    columns: int
    cells_dropped: bool  # a row ran past MAX_TABLE_COLUMNS


def write_docx(*, document: InlinedDocument, title: str) -> bytes:
    """Build a .docx from the inlined HTML; past the limits, end it early."""
    deadline = time.monotonic() + DEADLINE_SECONDS
    parser = lxml.html.HTMLParser(
        encoding="utf-8", no_network=True, huge_tree=False, remove_comments=True, remove_pis=True
    )
    uris = _SetAsideUris()
    html = uris.set_aside(text=_xml_safe(text=document.html))
    try:
        root = lxml.html.document_fromstring(html.encode("utf-8"), parser=parser)
    except ParserError:
        # lxml.html raises this for input with no element at all.
        root = None
    if root is not None:
        uris.put_back(root=root)
    word = docx.Document()
    _set_core_properties(
        word=word, title=_document_title(root=root) or _collapse(text=_xml_safe(text=title)).strip()
    )
    complete, cells_dropped = True, False
    if root is not None:
        writer = _Writer(word=word, deadline=deadline)
        complete = writer.write(root=root)
        cells_dropped = writer.cells_dropped
    if cells_dropped:
        word.add_paragraph(TABLE_CELLS_NOTE)
    # Past a resource limit, such as nesting deeper than 256, libxml2 stops
    # parsing and the tree holds only what came before.
    if not complete or parse_stopped(error_log=parser.error_log):
        word.add_paragraph(TRUNCATION_NOTE)
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


@dataclass
class _SetAsideUris:
    """Long ``data:`` URIs held out of the parse, each behind a short placeholder.

    libxml2 with huge_tree off stops reading after about 10 MB of large
    attribute or text tokens and drops the rest of the document, so a few
    inlined images would cut every block after them. The parser reads a
    placeholder from ``UriPlaceholders`` instead, and ``put_back`` restores
    each URI in the tree.
    """

    placeholders: UriPlaceholders = field(default_factory=UriPlaceholders)

    def set_aside(self, *, text: str) -> str:
        return _LONG_DATA_URI.sub(lambda match: self.placeholders.hold(uri=match.group()), text)

    def put_back(self, *, root: HtmlElement) -> None:
        """Swap each placeholder in a text, or in an attribute the writer reads, back for its URI.

        A placeholder anywhere else stays: nothing reads it, and the parser
        accepts attribute names, such as ``{a``, that lxml refuses to set.
        """
        if not self.placeholders.uris:
            return
        paths = [f"//@{name}[contains(., $nonce)]" for name in _READ_ATTRIBUTES]
        found = root.xpath(
            " | ".join([*paths, "//text()[contains(., $nonce)]"]), nonce=self.placeholders.nonce
        )
        for value in found:
            restored = self.placeholders.restore(text=value)
            owner = value.getparent()
            if value.is_attribute:
                owner.set(value.attrname, restored)
            elif value.is_tail:
                owner.tail = restored
            else:
                owner.text = restored


def _set_core_properties(*, word: WordDocument, title: str) -> None:
    """Set the title and replace what python-docx's template says about itself.

    The template names python-docx as author and description, dates itself
    2013, and carries a blank thumbnail that file browsers show as the
    preview.
    """
    properties = word.core_properties
    properties.title = title[:MAX_TITLE_LENGTH].rstrip()
    properties.author = ""
    properties.comments = ""
    now = datetime.now(UTC)
    properties.created = now
    properties.modified = now
    package_relationships = word.part.package.rels
    thumbnails = [key for key, value in package_relationships.items() if value.reltype == RT.THUMBNAIL]
    for relationship_id in thumbnails:
        del package_relationships[relationship_id]


class _Writer:
    def __init__(self, *, word: WordDocument, deadline: float) -> None:
        self._word = word
        self._deadline = deadline
        self._nodes = 0
        self._cells_left = MAX_TABLE_CELLS
        self._stopped = False
        # Set when a table limit leaves out a cell or a row.
        self.cells_dropped = False
        section = word.sections[0]
        self._page_width = section.page_width - section.left_margin - section.right_margin
        self._page_height = section.page_height - section.top_margin - section.bottom_margin
        self._section_properties = word.element.body.sectPr
        self._numbering = word.part.numbering_part.element
        self._next_num_id = max(int(value) for value in self._numbering.xpath("./w:num/@w:numId")) + 1
        self._abstract_num_ids: dict[str, int] = {}
        self._formatted_levels: dict[tuple[int, str], BaseOxmlElement] = {}
        self._link_ids: dict[str, str] = {}
        self._style_ids: dict[str, str] = {}
        self._next_relationship = 1

    def write(self, *, root: HtmlElement) -> bool:
        """Walk the tree into the document; False when a limit stopped the walk."""
        body = _Flow(container=self._word, width=self._page_width)
        stack = [_Visit(element=root, outer=_Context(flow=body))]
        while stack and not self._stopped:
            visit = stack.pop()
            if visit.leaving:
                self._leave(visit=visit)
                continue
            if visit.table_parts is not None:
                stack.extend(reversed(self._build_table(visit=visit)))
                continue
            following = visit.element.getnext() if visit.siblings else None
            if following is not None:
                stack.append(_Visit(element=following, outer=visit.outer, siblings=True))
            if visit.prepaid or self._spend():
                self._enter(visit=visit, stack=stack)
        return not self._stopped

    def _spend(self) -> bool:
        """Count one node; past the node budget or the deadline, stop the walk."""
        self._nodes += 1
        if self._nodes > MAX_NODES or time.monotonic() > self._deadline:
            self._stopped = True
        return not self._stopped

    def _enter(self, *, visit: _Visit, stack: list[_Visit]) -> None:
        element, outer = visit.element, visit.outer
        tag = element.tag if isinstance(element.tag, str) else ""
        text = element.text
        if tag in _SKIPPED_TAGS:
            self._write_tail(visit=visit)
            return
        if outer.depth >= MAX_DEPTH:
            # Past the depth limit, tags lose their meaning and only text flows.
            inner = outer
        elif tag == "br":
            self._line_break(context=outer)
            self._write_tail(visit=visit)
            return
        elif tag == "img":
            self._image(element=element, context=outer)
            self._write_tail(visit=visit)
            return
        elif tag == "table" and outer.table_depth < MAX_TABLE_NESTING:
            parts = self._table_parts(table=element)
            _close_paragraph(flow=outer.flow)
            # The stray content is written first, then the table is added.
            stray_context = replace(outer, depth=outer.depth + 1)
            stack.append(replace(visit, leaving=True))
            stack.append(replace(visit, table_parts=parts))
            stack.extend(_Visit(element=stray, outer=stray_context) for stray in reversed(parts.strays))
            return
        else:
            inner = replace(self._open(tag=tag, element=element, outer=outer), depth=outer.depth + 1)
            if tag == "pre" and text:
                # HTML drops one newline straight after <pre>.
                text = text[2:] if text.startswith("\r\n") else text[1:] if text[0] in "\r\n" else text
        if text:
            self._write_text(context=inner, raw=text)
        stack.append(replace(visit, leaving=True))
        first_child = next(iter(element), None)
        if first_child is not None:
            stack.append(_Visit(element=first_child, outer=inner, siblings=True))

    def _leave(self, *, visit: _Visit) -> None:
        tag = visit.element.tag
        if visit.outer.depth < MAX_DEPTH and (tag in _PARAGRAPH_TAGS or tag in _CONTAINER_TAGS):
            _close_paragraph(flow=visit.outer.flow)
        self._write_tail(visit=visit)

    def _write_tail(self, *, visit: _Visit) -> None:
        if visit.write_tail and visit.element.tail:
            self._write_text(context=visit.outer, raw=visit.element.tail)

    def _open(self, *, tag: str, element: HtmlElement, outer: _Context) -> _Context:
        """Apply what an element means on the way in; return its children's context."""
        flow = outer.flow
        if tag in _INLINE_STYLES:
            return replace(outer, run=replace(outer.run, **_INLINE_STYLES[tag]))
        if tag == "a":
            url = _link_target(href=element.get("href"))
            return replace(outer, run=replace(outer.run, link=_Link(url=url) if url else None))
        if tag in _HEADING_LEVELS:
            style = f"Heading {_HEADING_LEVELS[tag]}"
            if not self._take_item_paragraph(flow=flow, style=style):
                self._start_paragraph(flow=flow, style=style)
            return replace(outer, paragraph_style=style)
        if tag in _ITEM_REUSING_TAGS and flow.item_pending:
            return outer
        if tag == "pre":
            if not self._take_item_paragraph(flow=flow, style=None):
                self._start_paragraph(flow=flow, style=outer.paragraph_style)
            return replace(outer, pre=True, run=replace(outer.run, monospace=True))
        if tag in ("ul", "ol"):
            _close_paragraph(flow=flow)
            outer_level = outer.list_numbering.level if outer.list_numbering is not None else 0
            numbering = _list_numbering(
                element=element, kind=tag, level=min(outer_level + 1, MAX_LIST_LEVEL)
            )
            return replace(outer, list_numbering=numbering, paragraph_style=None)
        if tag == "li":
            # An item outside any list is a first-level bullet.
            numbering = outer.list_numbering
            level = numbering.level if numbering is not None else 1
            style = numbering.style if numbering is not None else _list_style(kind="ul", level=1)
            paragraph = self._start_paragraph(flow=flow, style=style)
            if numbering is not None:
                value = _leading_number(value=element.get("value"), ceiling=MAX_LIST_NUMBER)
                num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = 0
                num_pr.get_or_add_numId().val = self._item_num_id(
                    numbering=numbering, value=value if numbering.ordered else None
                )
            flow.item_pending = True
            return replace(outer, paragraph_style=_list_continue_style(level=level))
        if tag in _PARAGRAPH_TAGS:
            self._start_paragraph(flow=flow, style=outer.paragraph_style)
            return outer
        if tag == "blockquote":
            # A list item's first block keeps the item's numbered paragraph.
            if not flow.item_pending:
                _close_paragraph(flow=flow)
            return replace(outer, paragraph_style="Quote")
        if tag == "th":
            _close_paragraph(flow=flow)
            return replace(outer, run=replace(outer.run, bold=True))
        if tag in _CONTAINER_TAGS:
            _close_paragraph(flow=flow)
        return outer

    def _take_item_paragraph(self, *, flow: _Flow, style: str | None) -> bool:
        """Give a list item's first block the item's empty numbered paragraph.

        Without this, a heading or code block as an item's first block leaves
        the item's number alone on an empty line. False when no item waits.
        """
        if not flow.item_pending or flow.paragraph is None:
            return False
        if style is not None:
            self._set_style(paragraph=flow.paragraph, style=style)
        flow.item_pending = False
        return True

    def _start_paragraph(self, *, flow: _Flow, style: str | None) -> Paragraph:
        if flow.reusable is not None:
            paragraph = flow.reusable
            flow.reusable = None
        elif flow.container is self._word:
            # python-docx's add_paragraph searches the whole body for its
            # sectPr each time; inserting before the one we hold keeps a long
            # document linear.
            p = OxmlElement("w:p")
            self._section_properties.addprevious(p)
            paragraph = Paragraph(p, self._word)
        else:
            paragraph = flow.container.add_paragraph()
        if style:
            self._set_style(paragraph=paragraph, style=style)
        flow.paragraph = paragraph
        flow.open_link = None
        flow.at_space = True
        flow.item_pending = False
        return paragraph

    def _set_style(self, *, paragraph: Paragraph, style: str) -> None:
        """Set a paragraph style by name, resolving each name once.

        python-docx's style setter scans every style on each assignment.
        """
        style_id = self._style_ids.get(style)
        if style_id is None:
            style_id = self._word.styles[style].style_id
            self._style_ids[style] = style_id
        paragraph._p.style = style_id

    def _new_run(self, *, context: _Context) -> Run:
        """A run at the end of the current paragraph, inside a hyperlink when the style has one."""
        flow = context.flow
        paragraph = flow.paragraph
        if paragraph is None:
            paragraph = self._start_paragraph(flow=flow, style=context.paragraph_style)
        flow.item_pending = False
        link = context.run.link
        if link is None:
            flow.open_link = None
            run = paragraph.add_run()
        else:
            if flow.open_link is None or flow.open_link.link is not link:
                relationship_id = self._link_relationship(url=link.url)
                hyperlink = OxmlElement("w:hyperlink", attrs={qn("r:id"): relationship_id})
                paragraph._p.append(hyperlink)
                flow.open_link = _OpenLink(link=link, element=hyperlink)
            r = OxmlElement("w:r")
            flow.open_link.element.append(r)
            run = Run(r, paragraph)
        _format_run(run=run, style=context.run)
        return run

    def _link_relationship(self, *, url: str) -> str:
        """The id of the one external relationship for ``url``.

        python-docx's relate_to scans every relationship to reuse one and to
        pick a free id, which makes a document with many links quadratic.
        """
        relationship_id = self._link_ids.get(url)
        if relationship_id is None:
            relationships = self._word.part.rels
            while f"rId{self._next_relationship}" in relationships:
                self._next_relationship += 1
            relationship_id = f"rId{self._next_relationship}"
            relationships.add_relationship(RT.HYPERLINK, url, relationship_id, is_external=True)
            self._link_ids[url] = relationship_id
        return relationship_id

    def _write_text(self, *, context: _Context, raw: str) -> None:
        flow = context.flow
        text = _xml_safe(text=raw)
        if context.pre:
            text = self._fit_preformatted(text=text.replace("\r\n", "\n").replace("\r", "\n"))
        else:
            text = _collapse(text=text)
            if flow.paragraph is None or flow.at_space:
                text = text.lstrip(" ")
        if not text:
            return
        run = self._new_run(context=context)
        # python-docx turns "\n" into a line break and "\t" into a tab.
        run.text = text
        flow.at_space = not context.pre and text.endswith(" ")

    def _fit_preformatted(self, *, text: str) -> str:
        """Charge each line break and tab as a node; past the budget, cut the text and stop.

        python-docx writes each one as its own element, so a single text node
        holding millions of them would otherwise take minutes and gigabytes.
        """
        breaks = text.count("\n") + text.count("\t")
        room = MAX_NODES - self._nodes
        if breaks <= room:
            self._nodes += breaks
            return text
        self._nodes = MAX_NODES
        self._stopped = True
        cut = next(itertools.islice(_PRE_BREAKS.finditer(text), room, None)).start()
        return text[:cut]

    def _line_break(self, *, context: _Context) -> None:
        self._new_run(context=context).add_break()
        context.flow.at_space = True

    def _image(self, *, element: HtmlElement, context: _Context) -> None:
        alt = _collapse(text=_xml_safe(text=element.get("alt") or "")).strip()
        data = _image_bytes(src=element.get("src") or "")
        size = None
        if data is not None:
            size = _picture_size(data=data, max_width=context.flow.width, max_height=self._page_height)
        if data is None or size is None:
            self._write_text(context=context, raw=f"[image: {alt}]" if alt else "[image]")
            return
        run = self._new_run(context=context)
        shape = run.add_picture(io.BytesIO(data), width=Emu(size.width), height=Emu(size.height))
        shape._inline.docPr.set("descr", alt)
        context.flow.at_space = False

    def _item_num_id(self, *, numbering: _ListNumbering, value: int | None) -> int:
        """The w:num for a list's next item, and the number the item after it takes."""
        if value is not None:
            numbering.next_number = value
            numbering.num_id = None
        if numbering.num_id is None or numbering.reversed:
            numbering.num_id = self._numbering_clone(
                style=numbering.style,
                start=numbering.next_number if numbering.ordered else None,
                number_format=numbering.number_format,
            )
        step = -1 if numbering.reversed else 1
        numbering.next_number = min(max(numbering.next_number + step, 0), MAX_LIST_NUMBER)
        return numbering.num_id

    def _numbering_clone(self, *, style: str, start: int | None, number_format: str | None) -> int:
        """A new w:num on the style's abstract numbering, from ``start`` in ``number_format``."""
        abstract_id = self._abstract_num_ids.get(style)
        if abstract_id is None:
            style_num_id = self._word.styles[style].element.pPr.numPr.numId.val
            abstract_id = self._numbering.num_having_numId(style_num_id).abstractNumId.val
            self._abstract_num_ids[style] = abstract_id
        num_id = self._next_num_id
        self._next_num_id += 1
        num = CT_Num.new(num_id, abstract_id)
        if start is not None or number_format is not None:
            override = num.add_lvlOverride(ilvl=0)
            if start is not None:
                override.add_startOverride(start)
            if number_format is not None:
                override.append(self._formatted_level(abstract_id=abstract_id, number_format=number_format))
        # The numbering part ends with its w:num list, so appending keeps the
        # schema's order without the linear search add_num makes per list.
        self._numbering.append(num)
        return num_id

    def _formatted_level(self, *, abstract_id: int, number_format: str) -> BaseOxmlElement:
        """A copy of the abstract numbering's first level that shows ``number_format``.

        The copy keeps the level's indents and text, so only the format changes.
        """
        key = (abstract_id, number_format)
        level = self._formatted_levels.get(key)
        if level is None:
            [source] = self._numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="0"]')
            level = copy.deepcopy(source)
            level.find(qn("w:numFmt")).set(qn("w:val"), number_format)
            self._formatted_levels[key] = level
        return copy.deepcopy(level)

    def _table_parts(self, *, table: HtmlElement) -> _TableParts:
        """Sort a table's content into rows of td and th, captions, and stray content.

        As a browser does, rows and cells stay in the table even inside a
        wrapper such as form, cells with no row share one, and any other text
        or element is shown before the table. Each element sorted costs a node.
        """
        parts = _TableParts(rows=[], captions=[], strays=[])
        _add_stray_text(parts=parts, table=table, text=table.text)
        loose_row: list[HtmlElement] | None = None
        holding: set[HtmlElement] = set()
        stack: list[_TableItem] = []
        _queue_first_child(element=table, row=None, stack=stack)
        while stack:
            item = stack.pop()
            element = item.element
            if item.leaving:
                _add_stray_text(parts=parts, table=table, text=element.tail)
                continue
            following = element.getnext() if item.siblings else None
            if following is not None:
                stack.append(_TableItem(element=following, row=item.row, siblings=True))
            if not self._spend():
                break
            tag = element.tag if isinstance(element.tag, str) else ""
            row = item.row
            if tag in _TABLE_CELLS:
                if row is None:
                    if loose_row is None:
                        if len(parts.rows) >= self._cells_left:
                            self.cells_dropped = True
                            break
                        loose_row = []
                        parts.rows.append(loose_row)
                    row = loose_row
                if len(row) < MAX_TABLE_COLUMNS:
                    row.append(element)
                else:
                    self.cells_dropped = True
                _add_stray_text(parts=parts, table=table, text=element.tail)
            elif tag == "tr":
                loose_row = None
                if len(parts.rows) >= self._cells_left:
                    self.cells_dropped = True
                    break
                cells: list[HtmlElement] = []
                parts.rows.append(cells)
                _queue_children(element=element, row=cells, parts=parts, table=table, stack=stack)
            elif tag in _TABLE_ROW_GROUPS:
                loose_row = None
                _queue_children(element=element, row=None, parts=parts, table=table, stack=stack)
            elif tag == "caption":
                loose_row = None
                parts.captions.append(element)
                _add_stray_text(parts=parts, table=table, text=element.tail)
            elif tag in ("", "colgroup", "col"):
                # Column groups and non-element nodes show nothing themselves.
                _add_stray_text(parts=parts, table=table, text=element.tail)
            elif tag not in _SKIPPED_TAGS and tag != "table" and _holds_table_rows(element=element, holding=holding):
                # A wrapper around rows or cells: its rows join the table and
                # the rest of it is shown before the table.
                _queue_children(element=element, row=row, parts=parts, table=table, stack=stack)
            else:
                # Written whole before the table; its visit writes its tail.
                parts.strays.append(element)
        return parts

    def _build_table(self, *, visit: _Visit) -> list[_Visit]:
        """Add the captions and the table with its merges; return one visit per cell to fill."""
        parts, outer = visit.table_parts, visit.outer
        flow = outer.flow
        _close_paragraph(flow=flow)
        for caption in parts.captions:
            self._start_paragraph(flow=flow, style="Caption")
            self._write_text(context=outer, raw=caption.text_content())
            _close_paragraph(flow=flow)
        layout = _layout(rows=parts.rows)
        columns = layout.columns
        row_count = min(len(parts.rows), self._cells_left // columns) if columns else 0
        dropped_rows = any(placement.row >= row_count for placement in layout.placements)
        if layout.cells_dropped or dropped_rows:
            self.cells_dropped = True
        if row_count == 0:
            return []
        self._cells_left -= row_count * columns
        table = flow.container.add_table(rows=row_count, cols=columns)
        table.style = "Table Grid"
        # A table inside a cell is followed by the empty paragraph a cell must
        # end with; text after the table goes there.
        following = table._tbl.getnext()
        is_paragraph = following is not None and following.tag == qn("w:p")
        flow.reusable = Paragraph(following, flow.container) if is_paragraph else None
        row_elements = table._tbl.tr_lst
        grid = [row.tc_lst for row in row_elements]
        visits: list[_Visit] = []
        for placement in layout.placements:
            if placement.row >= row_count:
                continue
            rowspan = min(placement.rowspan, row_count - placement.row)
            for row in range(placement.row, placement.row + rowspan):
                tc = grid[row][placement.column]
                if placement.colspan > 1:
                    spanned = grid[row][placement.column : placement.column + placement.colspan]
                    tc.width = Emu(sum(cell.width or 0 for cell in spanned))
                    tc.grid_span = placement.colspan
                    for swallowed in spanned[1:]:
                        row_elements[row].remove(swallowed)
                if rowspan > 1:
                    tc.vMerge = ST_Merge.RESTART if row == placement.row else ST_Merge.CONTINUE
            cell = _Cell(grid[placement.row][placement.column], table)
            cell_flow = _Flow(container=cell, width=cell.width or flow.width, reusable=cell.paragraphs[0])
            visits.append(
                _Visit(
                    element=placement.element,
                    outer=_Context(
                        flow=cell_flow, depth=outer.depth + 2, table_depth=outer.table_depth + 1
                    ),
                    write_tail=False,
                    prepaid=True,
                )
            )
        return visits


def _queue_children(
    *,
    element: HtmlElement,
    row: list[HtmlElement] | None,
    parts: _TableParts,
    table: HtmlElement,
    stack: list[_TableItem],
) -> None:
    """Queue an element's children for sorting: its text goes now, its tail after them."""
    _add_stray_text(parts=parts, table=table, text=element.text)
    stack.append(_TableItem(element=element, row=None, leaving=True))
    _queue_first_child(element=element, row=row, stack=stack)


def _queue_first_child(*, element: HtmlElement, row: list[HtmlElement] | None, stack: list[_TableItem]) -> None:
    first_child = next(iter(element), None)
    if first_child is not None:
        stack.append(_TableItem(element=first_child, row=row, siblings=True))


def _add_stray_text(*, parts: _TableParts, table: HtmlElement, text: str | None) -> None:
    """Queue text found in a table outside its cells, unless it is only whitespace."""
    if text and _HTML_SPACE.sub("", text):
        span = table.makeelement("span")
        span.text = text
        parts.strays.append(span)


def _holds_table_rows(*, element: HtmlElement, holding: set[HtmlElement]) -> bool:
    """Whether a row, cell or row group sits anywhere inside ``element``.

    A match marks every element between it and ``element`` in ``holding``,
    so wrappers nested around the same rows are not searched again.
    """
    if element in holding:
        return True
    match = next(element.iter(*_TABLE_STRUCTURE), None)
    if match is None:
        return False
    ancestor = match.getparent()
    while ancestor is not None and ancestor is not element:
        holding.add(ancestor)
        ancestor = ancestor.getparent()
    return True


def _layout(*, rows: list[list[HtmlElement]]) -> _TableLayout:
    """Place cells on a grid of at most MAX_TABLE_COLUMNS columns, as HTML does.

    A cell starts at the first column no rowspan from above covers, and its
    colspan stops at the next covered column, so no two cells overlap.
    """
    covered = [0] * MAX_TABLE_COLUMNS
    placements: list[_CellPlacement] = []
    columns = 0
    cells_dropped = False
    for row_index, cells in enumerate(rows):
        column = 0
        rows_left = len(rows) - row_index
        for cell in cells:
            while column < MAX_TABLE_COLUMNS and covered[column]:
                column += 1
            if column >= MAX_TABLE_COLUMNS:
                cells_dropped = True
                break
            wanted = _span(value=cell.get("colspan"), limit=MAX_TABLE_COLUMNS, zero_means=1)
            colspan = 0
            while colspan < wanted and column + colspan < MAX_TABLE_COLUMNS and not covered[column + colspan]:
                colspan += 1
            rowspan = _span(value=cell.get("rowspan"), limit=rows_left, zero_means=rows_left)
            placements.append(
                _CellPlacement(element=cell, row=row_index, column=column, rowspan=rowspan, colspan=colspan)
            )
            for spanned in range(column, column + colspan):
                covered[spanned] = rowspan
            column += colspan
            columns = max(columns, column)
        covered = [max(0, left - 1) for left in covered]
    return _TableLayout(placements=placements, columns=columns, cells_dropped=cells_dropped)


def _span(*, value: str | None, limit: int, zero_means: int) -> int:
    span = _leading_number(value=value, ceiling=limit)
    if span is None:
        return 1
    return min(span or zero_means, limit)


def _leading_number(*, value: str | None, ceiling: int) -> int | None:
    """The number an attribute starts with, at most ``ceiling``; None with no leading digits."""
    match = _LEADING_DIGITS.match(value or "")
    if match is None:
        return None
    digits = match.group(1).lstrip("0")
    # int() refuses more than 4300 digits, and a number that long is past
    # any ceiling anyway.
    if len(digits) > len(str(ceiling)):
        return ceiling
    return min(int(digits or "0"), ceiling)


def _list_numbering(*, element: HtmlElement, kind: str, level: int) -> _ListNumbering:
    style = _list_style(kind=kind, level=level)
    if kind != "ol":
        return _ListNumbering(
            level=level, style=style, ordered=False, number_format=None, reversed=False, next_number=1
        )
    start = _leading_number(value=element.get("start"), ceiling=MAX_LIST_NUMBER)
    is_reversed = element.get("reversed") is not None
    if start is None:
        # A reversed list counts down to 1 from its number of items.
        start = min(sum(1 for child in element if child.tag == "li"), MAX_LIST_NUMBER) if is_reversed else 1
    return _ListNumbering(
        level=level,
        style=style,
        ordered=True,
        number_format=_NUMBER_FORMATS.get(element.get("type") or ""),
        reversed=is_reversed,
        next_number=start,
    )


def _list_style(*, kind: str, level: int) -> str:
    base = "List Number" if kind == "ol" else "List Bullet"
    return base if level == 1 else f"{base} {level}"


def _list_continue_style(*, level: int) -> str:
    return "List Continue" if level == 1 else f"List Continue {level}"


def _close_paragraph(*, flow: _Flow) -> None:
    flow.paragraph = None
    flow.open_link = None
    flow.item_pending = False


def _format_run(*, run: Run, style: _RunStyle) -> None:
    font = run.font
    if style.bold:
        font.bold = True
    if style.italic:
        font.italic = True
    if style.underline or style.link is not None:
        font.underline = True
    if style.strike:
        font.strike = True
    if style.monospace:
        font.name = MONOSPACE_FONT
    if style.subscript:
        font.subscript = True
    if style.superscript:
        font.superscript = True
    if style.highlight:
        font.highlight_color = WD_COLOR_INDEX.YELLOW
    if style.small:
        font.size = _SMALL_SIZE
    if style.link is not None:
        font.color.rgb = _LINK_COLOR


def _link_target(*, href: str | None) -> str | None:
    """The URL to link to, or None unless it is http, https or mailto."""
    candidate = _xml_safe(text=href or "").strip()
    try:
        parts = urlsplit(candidate)
    except ValueError:
        return None
    if parts.scheme not in LINK_SCHEMES:
        return None
    if parts.scheme == "mailto" and not parts.path:
        return None
    if parts.scheme != "mailto" and not parts.netloc:
        return None
    return quote(_LONE_PERCENT.sub("%25", parts.geturl()), safe=_URI_SAFE)


def _image_bytes(*, src: str) -> bytes | None:
    """The bytes of a base64 ``data:`` URI holding a PNG, JPEG or GIF, else None.

    The bytes decide the type, not the URI's media type.
    """
    if src[:5].lower() != "data:":
        return None
    header, comma, payload = src[5:].partition(",")
    if not comma or header.split(";")[-1].strip().lower() != "base64":
        return None
    try:
        data = base64.b64decode(_HTML_SPACE.sub("", payload), validate=True)
    except (binascii.Error, ValueError):
        return None
    if data.startswith((PNG_SIGNATURE, JPEG_SIGNATURE, *GIF_SIGNATURES)):
        return data
    return None


def _picture_size(*, data: bytes, max_width: int, max_height: int) -> _PictureSize | None:
    """The image's size in EMU at its own DPI, scaled down to fit the box.

    None when python-docx can't read the header: it reads PNG, JPEG with a
    JFIF or Exif marker, and GIF, and reads no pixels. None too for a header
    past ``_MAX_PNG_CHUNKS``, ``_MAX_JPEG_MARKERS`` or ``_MAX_EXIF_ENTRIES``,
    or with an Exif block shorter than its own header.
    """
    if not _header_is_bounded(data=data):
        return None
    try:
        image = Image.from_blob(data)
        px_width, px_height = image.px_width, image.px_height
        horz_dpi, vert_dpi = image.horz_dpi, image.vert_dpi
    except Exception:
        # Besides its own image errors, python-docx raises a bare Exception,
        # KeyError, struct.error or UnicodeDecodeError on a damaged header.
        return None
    if px_width <= 0 or px_height <= 0:
        return None
    width = px_width * _EMU_PER_INCH / (horz_dpi or _DEFAULT_DPI)
    height = px_height * _EMU_PER_INCH / (vert_dpi or _DEFAULT_DPI)
    scale = min(1.0, max_width / width, max_height / height)
    return _PictureSize(width=max(1, int(width * scale)), height=max(1, int(height * scale)))


def _header_is_bounded(*, data: bytes) -> bool:
    """Whether python-docx reads this image's header without a blowup.

    It picks the parser in python-docx's order: the PNG signature first,
    then a JFIF or Exif tag at offset 6, and only then the GIF signature.
    So bytes with a GIF signature and either tag are read as a JPEG.
    """
    if data.startswith(PNG_SIGNATURE):
        return _png_chunk_count(data=data) <= _MAX_PNG_CHUNKS
    if data[6:10] in (b"JFIF", b"Exif"):
        return _jpeg_header_is_bounded(data=data)
    # python-docx reads a GIF only for its screen size, and raises at once
    # on a JPEG with neither tag.
    return True


def _png_chunk_count(*, data: bytes) -> int:
    """The chunks python-docx reads, up to IEND; it stops counting past the cap."""
    count, offset = 0, len(PNG_SIGNATURE)
    while offset + 8 <= len(data) and count <= _MAX_PNG_CHUNKS:
        count += 1
        if data[offset + 4 : offset + 8] == b"IEND":
            break
        offset += 12 + int.from_bytes(data[offset : offset + 4], "big")
    return count


def _jpeg_header_is_bounded(*, data: bytes) -> bool:
    """Whether python-docx reads the markers before a JPEG's scan within the caps.

    It finds each marker as python-docx does: the next 0xFF after the last
    segment, past any 0xFF fill, where an 0xFF 0x00 pair is no marker. It
    stops as soon as a cap is passed, so its own cost stays small.
    """
    markers, entries, position = 0, 0, 0
    while (ff := data.find(b"\xff", position)) != -1:
        position = ff + 1
        while position < len(data) and data[position] == 0xFF:
            position += 1
        if position >= len(data) or data[position] in (_JPEG_EOI, _JPEG_SOS):
            break
        marker = data[position]
        position += 1
        markers += 1
        if markers > _MAX_JPEG_MARKERS:
            return False
        if marker == 0x00 or marker in JPEG_STANDALONE_MARKERS:
            continue
        length = int.from_bytes(data[position : position + 2], "big")
        if marker == _JPEG_APP1 and data[position + 2 : position + 8] == b"Exif\x00\x00":
            # python-docx reads an Exif block shorter than its own header to
            # the end of the file.
            if length < 8:
                return False
            entries += _ifd0_entry_count(tiff=data[position + 8 : position + length])
            if entries > _MAX_EXIF_ENTRIES:
                return False
        position += length
    return True


def _ifd0_entry_count(*, tiff: bytes) -> int:
    """How many entries the first directory of a TIFF block declares."""
    order = "big" if tiff[:2] == b"MM" else "little"
    directory = int.from_bytes(tiff[4:8], order)
    return int.from_bytes(tiff[directory : directory + 2], order)


def _document_title(*, root: HtmlElement | None) -> str:
    if root is None:
        return ""
    head = root.find("head")
    title = head.find("title") if head is not None else None
    if title is None:
        return ""
    return _collapse(text=_xml_safe(text=title.text_content())).strip()


def _xml_safe(*, text: str) -> str:
    return XML_UNSAFE.sub("", text)


def _collapse(*, text: str) -> str:
    return _HTML_SPACE.sub(" ", text)
