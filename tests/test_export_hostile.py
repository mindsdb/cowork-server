"""Hostile references through the whole export path: read, render, inline, write.

Each YAML case under ``tests/fixtures/export_hostile/`` runs as an HTML source
in Word and PDF modes, and as a Markdown source in HTML, Word and PDF modes.
The artifact folder and another partition sit side by side under ``tmp_path``.
Inside the ``no_fetch()`` guard the source is read through the pinned folder,
rendered when it is Markdown, inlined, and written to a .docx in Word mode. A
127.0.0.1 recording server, started before the guard, is the outside target,
so it also sees a connection native code makes. Afterwards the guard and the
server have seen nothing, no output holds the other partition's sentinel
bytes, and every expected reference is listed in ``omitted`` with its reason.

Case fields:
  id               names the case
  html, markdown   the source, as an HTML file and as a Markdown file
  omitted          [reason, reference] pairs that ``omitted`` must hold, as the
                   inliner records them: control characters stripped, not decoded
  windows_omitted  replaces ``omitted`` on Windows
  files            extra files written into the artifact folder
  text             text the inlined HTML, and in Word mode the .docx, must hold
  present          strings the inlined HTML must hold
  kept_image       true when the artifact's own chart.png must be inlined

Placeholders, filled at run time in the sources, the files and the pairs:
  {server_url}              the recording server
  {sentinel_path}           the other partition's secret.png, as an absolute path
  {sentinel_file_url}       the same file as a file:// URL
  {sentinel_text_file_url}  the other partition's secret.txt as a file:// URL

No case links to {server_url}, so no attribute, style or script in the
output, and no relationship in the .docx, may hold it. Body text may: a
Markdown source turns a DOCTYPE internal subset into an inert paragraph.
"""
from __future__ import annotations

import base64
import io
import os
import re
import struct
import zipfile
import zlib
from collections import Counter
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

import docx
import lxml.html
import pytest
import yaml
from docx.document import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from lxml import etree

from cowork.common.paths import PinnedDir, pinned_dir
from cowork.schemas.artifact_export import (
    OMITTED_REFERENCE_MAX_CHARS,
    InlinedDocument,
    InlineLimits,
    InlineMode,
)
from cowork.services.artifact_files import read_artifact_file
from cowork.services.artifact_inline import inline_artifact_html
from cowork.services.docx_writer import write_docx
from cowork.services.markdown_render import render_markdown
from tests._no_fetch import RecordingServer, no_fetch, recording_server

FIXTURES = Path(__file__).parent / "fixtures" / "export_hostile"
CASE_FILES = ("outside.yaml", "partition.yaml", "css.yaml", "images.yaml", "parsers.yaml")

SENTINEL = b"SENTINEL-B"
HOUSEKEEPING_MARKER = b"HOUSEKEEPING-MARKER"
FORBIDDEN_BYTES = (SENTINEL, HOUSEKEEPING_MARKER)

# The read bounds and inline budgets the export route uses.
HTML_READ_BOUND = 10 * 1024 * 1024
MARKDOWN_READ_BOUND = 512 * 1024
INLINE_BUDGETS = {
    InlineMode.HTML: 25 * 1024 * 1024,
    InlineMode.WORD: 25 * 1024 * 1024,
    InlineMode.PDF: 5 * 1024 * 1024,
}

_SAFE_XML = etree.XMLParser(
    resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False
)
_DATA_URI_PAYLOAD = re.compile(r"base64,([A-Za-z0-9+/=]+)")


def _png(*, width: int, height: int, text: bytes = b"") -> bytes:
    """A decodable grayscale PNG, optionally carrying *text* in a tEXt chunk."""

    def chunk(*, kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))

    scanlines = (b"\x00" + b"\x80" * width) * height
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(kind=b"IHDR", body=struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0))
        + (chunk(kind=b"tEXt", body=b"note\x00" + text) if text else b"")
        + chunk(kind=b"IDAT", body=zlib.compress(scanlines))
        + chunk(kind=b"IEND", body=b"")
    )


CHART_PNG = _png(width=2, height=2)
CHART_DATA_URI = "data:image/png;base64," + base64.b64encode(CHART_PNG).decode("ascii")

# The artifact's housekeeping files hold bytes that would survive a read: a
# PNG where an img would inline it, CSS where an @import would splice it.
ARTIFACT_FILES = {
    "chart.png": CHART_PNG,
    "metadata.json": _png(width=1, height=1, text=HOUSEKEEPING_MARKER),
    "sub/metadata.json": b'.housekeeping::after{content:"' + HOUSEKEEPING_MARKER + b'"}',
    ".revisions/old.png": _png(width=1, height=1, text=HOUSEKEEPING_MARKER),
}
# Another person's partition, beside the artifact folder. Each file would
# carry the sentinel into the output if it were ever read.
OTHER_FILES = {
    "secret.png": _png(width=1, height=1, text=SENTINEL),
    "secret.txt": SENTINEL,
    "secret.css": b'body{background:url(secret.png)}.leak::after{content:"' + SENTINEL + b'"}',
    "secret.svg": b'<svg xmlns="http://www.w3.org/2000/svg"><text>' + SENTINEL + b"</text></svg>",
}


class Source(StrEnum):
    HTML = "html"
    MARKDOWN = "markdown"

    @property
    def file_name(self) -> str:
        return "report.html" if self is Source.HTML else "report.md"


SOURCE_MODES = (
    (Source.HTML, InlineMode.WORD),
    (Source.HTML, InlineMode.PDF),
    (Source.MARKDOWN, InlineMode.HTML),
    (Source.MARKDOWN, InlineMode.WORD),
    (Source.MARKDOWN, InlineMode.PDF),
)


@dataclass(frozen=True)
class Partitions:
    art: Path
    other: Path


@dataclass(frozen=True)
class Export:
    document: InlinedDocument
    package: bytes | None  # the .docx, in Word mode only


def _write_tree(*, base: Path, files: dict[str, bytes]) -> None:
    for name, data in files.items():
        path = base / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)


def _partitions(*, tmp_path: Path) -> Partitions:
    partitions = Partitions(art=tmp_path / "art", other=tmp_path / "other")
    _write_tree(base=partitions.art, files=ARTIFACT_FILES)
    _write_tree(base=partitions.other, files=OTHER_FILES)
    return partitions


def _placeholders(*, server: RecordingServer, partitions: Partitions) -> dict[str, str]:
    return {
        "server_url": server.url,
        "sentinel_path": str(partitions.other / "secret.png"),
        "sentinel_file_url": (partitions.other / "secret.png").as_uri(),
        "sentinel_text_file_url": (partitions.other / "secret.txt").as_uri(),
    }


def _fill(*, text: str, values: dict[str, str]) -> str:
    for name, value in values.items():
        text = text.replace("{" + name + "}", value)
    return text


def _export(*, folder: PinnedDir, source: Source, mode: InlineMode) -> Export:
    """The export path, from the source file in the pinned folder to the output."""
    if source is Source.MARKDOWN:
        text = read_artifact_file(
            folder=folder, parts=(source.file_name,), max_bytes=MARKDOWN_READ_BOUND
        )
        html = render_markdown(text=text.decode("utf-8")).encode("utf-8")
    else:
        html = read_artifact_file(folder=folder, parts=(source.file_name,), max_bytes=HTML_READ_BOUND)
    document = inline_artifact_html(
        html=html,
        folder=folder,
        base_parts=(),
        mode=mode,
        limits=InlineLimits(max_total_bytes=INLINE_BUDGETS[mode]),
    )
    package = write_docx(document=document, title="report") if mode is InlineMode.WORD else None
    return Export(document=document, package=package)


def _guarded_export(
    *, partitions: Partitions, server: RecordingServer, source: Source, mode: InlineMode
) -> Export:
    """Run the export inside the guard; neither the guard nor the server may see a call."""
    with pinned_dir(partitions.art) as folder:
        with no_fetch() as recorder:
            export = _export(folder=folder, source=source, mode=mode)
    assert recorder.calls == []
    assert server.hits() == []
    return export


def _members(*, package: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _assert_html_holds_nothing_forbidden(*, html: str, server_url: str) -> None:
    parser = lxml.html.HTMLParser(encoding="utf-8", no_network=True, huge_tree=False)
    root = lxml.html.document_fromstring(html.encode("utf-8"), parser=parser)
    for element in root.iter():
        for name, value in element.attrib.items():
            assert server_url not in value, f"{element.tag}@{name}"
        if element.tag in ("style", "script"):
            assert server_url not in (element.text or ""), element.tag
    for forbidden in FORBIDDEN_BYTES:
        assert forbidden.decode("ascii") not in html
    for payload in _DATA_URI_PAYLOAD.findall(html):
        decoded = base64.b64decode(payload)
        for forbidden in FORBIDDEN_BYTES:
            assert forbidden not in decoded


def _assert_package_is_clean(*, package: bytes, server_url: str) -> WordDocument:
    """Reload the .docx and check every member: nothing forbidden, links only outward."""
    document = docx.Document(io.BytesIO(package))
    members = _members(package=package)
    external = 0
    for name, body in members.items():
        for forbidden in FORBIDDEN_BYTES:
            assert forbidden not in body, f"{forbidden!r} in {name}"
        if name.endswith(".rels"):
            assert server_url.encode("ascii") not in body, name
            for relationship in etree.fromstring(body, _SAFE_XML):
                if relationship.get("TargetMode") == "External":
                    assert relationship.get("Type") == RT.HYPERLINK, relationship.get("Target")
                    external += 1
    assert sum(body.count(b'TargetMode="External"') for body in members.values()) == external
    return document


def _docx_text(*, document: WordDocument) -> str:
    return "".join(node.text or "" for node in document.element.body.iter(qn("w:t")))


def _expected_omitted(*, pairs: list[list[str]], values: dict[str, str]) -> Counter:
    return Counter(
        (reason, _fill(text=reference, values=values)[:OMITTED_REFERENCE_MAX_CHARS])
        for reason, reference in pairs
    )


def _cases() -> list:
    params = []
    for file_name in CASE_FILES:
        for case in yaml.safe_load((FIXTURES / file_name).read_text(encoding="utf-8")):
            for source, mode in SOURCE_MODES:
                params.append(
                    pytest.param(case, source, mode, id=f"{case['id']}-{source}-{mode}")
                )
    return params


@pytest.mark.parametrize(("case", "source", "mode"), _cases())
def test_hostile_case_exports_only_the_artifacts_own_content(case, source, mode, tmp_path):
    partitions = _partitions(tmp_path=tmp_path)
    with recording_server() as server:
        values = _placeholders(server=server, partitions=partitions)
        extra = {
            name: _fill(text=text, values=values).encode("utf-8")
            for name, text in case.get("files", {}).items()
        }
        _write_tree(base=partitions.art, files=extra)
        source_text = _fill(text=case[source.value], values=values)
        (partitions.art / source.file_name).write_bytes(source_text.encode("utf-8"))
        export = _guarded_export(partitions=partitions, server=server, source=source, mode=mode)

    html = export.document.html
    _assert_html_holds_nothing_forbidden(html=html, server_url=server.url)
    pairs = case["omitted"]
    if os.name == "nt" and "windows_omitted" in case:
        pairs = case["windows_omitted"]
    recorded = Counter((entry.reason.value, entry.reference) for entry in export.document.omitted)
    assert _expected_omitted(pairs=pairs, values=values) - recorded == Counter(), recorded
    for text in case.get("present", []):
        assert text in html
    if "text" in case:
        assert case["text"] in html
    if case.get("kept_image"):
        assert CHART_DATA_URI in html
    else:
        assert "data:image/" not in html

    if mode is not InlineMode.WORD:
        assert export.package is None
        return
    document = _assert_package_is_clean(package=export.package, server_url=server.url)
    if "text" in case:
        assert case["text"] in _docx_text(document=document)
    media = [
        body
        for name, body in _members(package=export.package).items()
        if name.startswith("word/media/")
    ]
    assert (CHART_PNG in media) is bool(case.get("kept_image"))


def test_the_manual_check_report_exports_to_word_with_only_its_own_content(tmp_path):
    """The report from the manual check: the chart and one link stay, two references go."""
    report = yaml.safe_load((FIXTURES / "report.yaml").read_text(encoding="utf-8"))
    partitions = _partitions(tmp_path=tmp_path)
    with recording_server() as server:
        values = _placeholders(server=server, partitions=partitions)
        source = _fill(text=report["html"], values=values)
        (partitions.art / Source.HTML.file_name).write_bytes(source.encode("utf-8"))
        export = _guarded_export(
            partitions=partitions, server=server, source=Source.HTML, mode=InlineMode.WORD
        )

    assert [(entry.reason.value, entry.reference) for entry in export.document.omitted] == [
        (reason, _fill(text=reference, values=values)) for reason, reference in report["omitted"]
    ]
    document = _assert_package_is_clean(package=export.package, server_url=server.url)
    assert len(document.inline_shapes) == 1
    members = _members(package=export.package)
    assert [body for name, body in members.items() if name.startswith("word/media/")] == [CHART_PNG]
    assert not [name for name, body in members.items() if SENTINEL in body]
    relationships = members["word/_rels/document.xml.rels"]
    assert relationships.count(b'TargetMode="External"') == 1
    external = [
        relationship
        for relationship in etree.fromstring(relationships, _SAFE_XML)
        if relationship.get("TargetMode") == "External"
    ]
    assert [(link.get("Type"), link.get("Target")) for link in external] == [
        (RT.HYPERLINK, report["link"])
    ]
