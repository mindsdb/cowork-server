"""POST /api/v1/artifacts/workspace/{project_ref}/{artifact_id}/export.

Local tenancy goes over HTTP, with artifact resolution pointed at a folder the
test builds, as the workspace identity-gate tests do. Org tenancy calls the
handler with a scoped session over real project rows: the app wires the org
principal middleware only when the process starts in org mode, so a request
would run under the local scope and prove nothing about who may export.
"""
from __future__ import annotations

import codecs
import errno
import io
import json
import os
import struct
import threading
import time
import zlib
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from urllib.parse import quote
from uuid import uuid4

import anyio
import docx
import httpx
import pytest
import yaml
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlmodel import Session

from cowork.api.v1.endpoints import artifact_workspace as workspace_ep
from cowork.common.settings.app_settings import get_app_settings, get_export_pdf_settings
from cowork.db.scoped import ScopedSession, TenantScope
from cowork.db.session import get_engine
from cowork.schemas.artifact_export import ExportPdfSettings, ExportRequest, ExportResult, RenderPdfRequest
from cowork.services import artifact_export
from cowork.services import artifact_ownership as ownership
from cowork.services import artifacts as artifacts_service
from cowork.services.artifact_export import ExportRefused, ExportSource, export_slot
from cowork.services.artifacts import ProjectArtifacts
from cowork.services.markdown_render import MAX_FENCE_LENGTH
from cowork.services.render_pdf_client import RenderServicePdf
from test_artifact_ownership import legacy_root, make_conversation, make_project, project_root

FIXTURES = Path(__file__).parent / "fixtures" / "artifact_export"
ROUTE_CASES = yaml.safe_load((FIXTURES / "route_cases.yaml").read_text(encoding="utf-8"))["cases"]
ARTIFACT_ID = "0123456789abcdef0123456789abcdef"
EXPORT_URL = f"/api/v1/artifacts/workspace/local/{ARTIFACT_ID}/export"
MIB = 1024 * 1024

pytestmark = pytest.mark.usefixtures("cleanup_tmp_projects")


def _snapshot(folder: Path) -> dict[str, object]:
    """Every entry below *folder*: a file's bytes, or a symlink's target."""
    entries: dict[str, object] = {}
    for directory, dirnames, filenames in os.walk(folder):
        for name in [*dirnames, *filenames]:
            path = Path(directory) / name
            key = path.relative_to(folder).as_posix()
            if path.is_symlink():
                entries[key] = ("symlink", os.readlink(path))
            elif path.is_file():
                entries[key] = path.read_bytes()
    return entries


def _metadata(folder: Path) -> dict:
    return json.loads((folder / "metadata.json").read_text(encoding="utf-8"))


# ── local tenancy ────────────────────────────────────────────────────────────


@pytest.fixture
def client(monkeypatch):
    from cowork.server import create_app

    monkeypatch.setenv("COWORK_TENANCY_MODE", "local")
    monkeypatch.delenv("COWORK_TRUST_IDENTITY_HEADERS", raising=False)
    get_app_settings.cache_clear()
    yield TestClient(create_app())
    get_app_settings.cache_clear()


@pytest.fixture
def local_artifact(tmp_path, monkeypatch):
    """Build one artifact folder that every export request resolves to."""

    def build(
        files: dict[str, str | bytes], *, primary: str | None = None, artifact_type: str = "document"
    ) -> Path:
        project = tmp_path / "project"
        base = project / ".anton" / "artifacts"
        folder = base / "report"
        folder.mkdir(parents=True)
        for name, content in files.items():
            path = folder / name
            path.parent.mkdir(parents=True, exist_ok=True)
            if isinstance(content, bytes):
                path.write_bytes(content)
            else:
                path.write_text(content, encoding="utf-8")
        metadata = {"id": ARTIFACT_ID, "slug": "report", "name": "Report", "type": artifact_type}
        if primary is not None:
            metadata["primary"] = primary
        (folder / "metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
        source = ProjectArtifacts(
            base=base,
            project_id=None,
            project_name="project",
            trusted_anchor=project,
            root_parts=(".anton", "artifacts"),
        )
        monkeypatch.setattr(
            workspace_ep,
            "review_artifact_for_request",
            lambda *_args: (source, folder, _metadata(folder), True),
        )
        return folder

    return build


def _assert_reloads(path: Path, export_format: str) -> None:
    if export_format == "docx":
        docx.Document(io.BytesIO(path.read_bytes()))
    else:
        text = path.read_text(encoding="utf-8")
        assert text[:15].lower() == "<!doctype html>"
        assert f"<title>{Path(path.name).stem}</title>" in text


@pytest.mark.parametrize("case", ROUTE_CASES, ids=[case["id"] for case in ROUTE_CASES])
def test_route_cases(client, local_artifact, case):
    folder = local_artifact(
        case["files"], primary=case["primary"], artifact_type=case.get("type", "document")
    )
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": case["format"]})

    assert response.status_code == case["status"], response.text
    after = _snapshot(folder)
    if case["status"] != 200:
        assert after == before
        return
    written = case["written"]
    assert response.json() == {
        "filename": Path(written).name,
        "draftUrl": f"/api/v1/artifacts/drafts/local/{ARTIFACT_ID}/{quote(written)}",
        "serveUrl": f"/api/v1/artifacts/serve/project/report/{quote(written)}",
        "omittedCount": 0,
        "scriptsDropped": case.get("scripts_dropped", False),
    }
    assert set(after) - set(before) == {written}
    assert {key: after[key] for key in before} == before
    _assert_reloads(folder / written, case["format"])


@pytest.mark.parametrize(
    "body",
    [
        {"format": "docx", "path": "/etc/passwd"},
        {"path": "report.md", "format": "html"},
        {"format": "docx", "filename": "report.docx"},
        {"format": "rtf"},
        {"format": "DOCX"},
        {},
    ],
)
def test_any_body_but_one_format_answers_422_and_writes_nothing(
    client, local_artifact, monkeypatch, body
):
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    before = _snapshot(folder)
    monkeypatch.setattr(
        workspace_ep, "export_primary_file", Mock(side_effect=AssertionError("no export"))
    )

    response = client.post(EXPORT_URL, json=body)

    assert response.status_code == 422, response.text
    assert _snapshot(folder) == before


def test_a_second_export_takes_the_next_free_name(client, local_artifact):
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")

    first = client.post(EXPORT_URL, json={"format": "docx"})
    original = (folder / "report.docx").read_bytes()
    second = client.post(EXPORT_URL, json={"format": "docx"})
    third = client.post(EXPORT_URL, json={"format": "docx"})

    assert [r.json()["filename"] for r in (first, second, third)] == [
        "report.docx",
        "report (2).docx",
        "report (3).docx",
    ]
    assert second.json()["draftUrl"].endswith("/report%20%282%29.docx")
    assert (folder / "report.docx").read_bytes() == original
    docx.Document(io.BytesIO((folder / "report (2).docx").read_bytes()))


@pytest.mark.parametrize("dangling", [False, True])
def test_a_symlink_at_the_output_name_is_never_written_through(
    client, local_artifact, tmp_path, dangling
):
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    victim = tmp_path / "victim.docx"
    if not dangling:
        victim.write_bytes(b"victim bytes")
    (folder / "report.docx").symlink_to(victim)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 200, response.text
    assert response.json()["filename"] == "report (2).docx"
    assert (folder / "report.docx").is_symlink()
    assert os.readlink(folder / "report.docx") == str(victim)
    if dangling:
        assert not victim.exists()
    else:
        assert victim.read_bytes() == b"victim bytes"


def test_a_hard_link_at_the_output_name_is_never_written_through(client, local_artifact, tmp_path):
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    victim = tmp_path / "victim.html"
    victim.write_bytes(b"victim bytes")
    os.link(victim, folder / "report.html")

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text
    assert response.json()["filename"] == "report (2).html"
    assert victim.read_bytes() == b"victim bytes"
    assert (folder / "report.html").stat().st_ino == victim.stat().st_ino


@pytest.mark.parametrize("planted", ["symlink", "hard-link", "file"])
def test_an_entry_planted_at_the_temporary_name_is_never_written_through(
    client, local_artifact, tmp_path, monkeypatch, planted
):
    """The temporary file is created with O_EXCL and O_NOFOLLOW, so an entry
    already at its name fails the export instead of taking its bytes."""
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    monkeypatch.setattr(
        artifact_export, "secrets", SimpleNamespace(token_hex=lambda nbytes: "0" * (2 * nbytes))
    )
    temporary = folder / f".cowork-export-{'0' * 24}.tmp"
    victim = tmp_path / "victim"
    victim.write_bytes(b"victim bytes")
    if planted == "symlink":
        temporary.symlink_to(victim)
    elif planted == "hard-link":
        os.link(victim, temporary)
    else:
        temporary.write_bytes(b"planted bytes")
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 500, response.text
    assert _snapshot(folder) == before
    assert victim.read_bytes() == b"victim bytes"


def _no_hard_links(code: int):
    """Make link() fail the way a volume without hard links fails it."""

    def refuse(*_args, **_kwargs):
        raise OSError(code, os.strerror(code))

    return refuse


@pytest.mark.parametrize(
    "code", [errno.ENOTSUP, errno.EOPNOTSUPP, errno.EPERM], ids=["ENOTSUP", "EOPNOTSUPP", "EPERM"]
)
@pytest.mark.parametrize("planted", ["symlink", "dangling-symlink", "file"])
def test_a_volume_without_hard_links_still_exports_and_replaces_nothing(
    client, local_artifact, tmp_path, monkeypatch, code, planted
):
    """exFAT and FAT32 refuse link(): the export writes the free name directly,
    and an entry planted at the first name is still never written through."""
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    victim = tmp_path / "victim.docx"
    if planted == "file":
        (folder / "report.docx").write_bytes(b"old bytes")
    else:
        if planted == "symlink":
            victim.write_bytes(b"victim bytes")
        (folder / "report.docx").symlink_to(victim)
    before = _snapshot(folder)
    monkeypatch.setattr(os, "link", _no_hard_links(code))

    first = client.post(EXPORT_URL, json={"format": "docx"})
    second = client.post(EXPORT_URL, json={"format": "html"})

    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert first.json()["filename"] == "report (2).docx"
    assert second.json()["filename"] == "report.html"
    after = _snapshot(folder)
    assert set(after) - set(before) == {"report (2).docx", "report.html"}
    assert {key: after[key] for key in before} == before
    if planted == "symlink":
        assert victim.read_bytes() == b"victim bytes"
    elif planted == "dangling-symlink":
        assert not victim.exists()
    docx.Document(io.BytesIO((folder / "report (2).docx").read_bytes()))
    _assert_reloads(folder / "report.html", "html")


def test_a_failed_write_without_hard_links_leaves_nothing(client, local_artifact, monkeypatch):
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    before = _snapshot(folder)
    monkeypatch.setattr(os, "link", _no_hard_links(errno.ENOTSUP))
    written = artifact_export._write_all
    calls: list[int] = []

    def full_disk(fd, data):
        # The temporary file is written first, then the output name.
        calls.append(fd)
        if len(calls) == 2:
            raise OSError(errno.ENOSPC, "No space left on device")
        written(fd, data)

    monkeypatch.setattr(artifact_export, "_write_all", full_disk)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 500, response.text
    assert len(calls) == 2
    assert _snapshot(folder) == before


@pytest.mark.parametrize("stem", ["a" * 250, "報" * 84], ids=["ascii-250-bytes", "cjk-252-bytes"])
def test_a_source_name_near_the_name_limit_still_exports(client, local_artifact, stem):
    """The stem plus " (2)" and the new extension would pass the 255-byte
    name limit, so the stem is cut, the same way for every output name."""
    folder = local_artifact({f"{stem}.md": "# R\n"}, primary=f"{stem}.md")

    names = []
    for _ in range(2):
        response = client.post(EXPORT_URL, json={"format": "docx"})
        assert response.status_code == 200, response.text
        names.append(response.json()["filename"])

    assert all(len(os.fsencode(name)) <= 255 for name in names)
    cut = names[0].removesuffix(".docx")
    assert names[1] == f"{cut} (2).docx"
    assert stem.startswith(cut) and len(cut) < len(stem)
    for name in names:
        docx.Document(io.BytesIO((folder / name).read_bytes()))


def test_an_html_export_of_markdown_writes_one_document_beside_it(client, local_artifact):
    source = "# Quarterly report\n\nRevenue *grew*.\n\n<script>alert(1)</script>\n"
    folder = local_artifact({"report.md": source}, primary="report.md")

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text
    assert (folder / "report.md").read_text(encoding="utf-8") == source
    written = (folder / "report.html").read_text(encoding="utf-8")
    assert written[:15].lower() == "<!doctype html>"
    assert written.count("<html") == 1 and written.count("<body") == 1
    assert "<title>report</title>" in written
    assert "Quarterly report</h1>" in written
    assert "<em>grew</em>" in written


_MARKED_SOURCE = "# Título\n\nCuerpo del informe.\n"


@pytest.mark.parametrize(
    "encoded",
    [
        codecs.BOM_UTF8 + _MARKED_SOURCE.encode("utf-8"),
        codecs.BOM_UTF16_LE + _MARKED_SOURCE.encode("utf-16-le"),
        codecs.BOM_UTF16_BE + _MARKED_SOURCE.encode("utf-16-be"),
    ],
    ids=["utf-8", "utf-16-le", "utf-16-be"],
)
def test_a_source_with_a_byte_order_mark_keeps_its_first_heading(client, local_artifact, encoded):
    """Windows editors start a file with a byte order mark. It is no text, so
    the first line stays a heading in Word and in HTML."""
    folder = local_artifact({"report.md": encoded}, primary="report.md")

    word = client.post(EXPORT_URL, json={"format": "docx"})
    page = client.post(EXPORT_URL, json={"format": "html"})

    assert word.status_code == 200, word.text
    assert page.status_code == 200, page.text
    paragraphs = docx.Document(io.BytesIO((folder / "report.docx").read_bytes())).paragraphs
    assert [(p.style.name, p.text) for p in paragraphs[:2]] == [
        ("Heading 1", "Título"),
        ("Normal", "Cuerpo del informe."),
    ]
    written = (folder / "report.html").read_text(encoding="utf-8")
    assert "Título</h1>" in written
    assert "\ufeff" not in written and "\ufffd" not in written


def test_a_text_source_with_a_byte_order_mark_exports_without_it(client, local_artifact):
    """A plain text file is read the same way as Markdown, so its byte order
    mark never reaches the exported page."""
    encoded = codecs.BOM_UTF8 + _MARKED_SOURCE.encode("utf-8")
    folder = local_artifact({"notes.txt": encoded}, primary="notes.txt")

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text
    written = (folder / "notes.html").read_text(encoding="utf-8")
    assert "<pre># Título" in written
    assert "\ufeff" not in written


def test_a_large_html_export_is_written_whole(client, local_artifact):
    """Eight 1.4 MiB images fit the 25 MiB budget and make a document over
    10 MB, past where lxml's HTML parser stops unless its limits are lifted.
    Every image and the text after them reach the file."""
    images = {f"chart{i}.png": padded_png(int(1.4 * MIB)) for i in range(8)}
    markdown = "\n\n".join(f"![chart]({name})" for name in images) + "\n\nTHE-END\n"
    folder = local_artifact({"report.md": markdown, **images}, primary="report.md")

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text
    assert response.json()["omittedCount"] == 0
    written = (folder / "report.html").read_text(encoding="utf-8")
    assert len(written) > 10_000_000
    assert written.count("data:image/png;base64,") == 8
    assert "THE-END" in written


@pytest.mark.parametrize(
    ("name", "size"),
    [
        ("page.html", 10 * MIB + 1),
        ("page.htm", 10 * MIB + 1),
        ("report.md", 512 * 1024 + 1),
        ("report.markdown", 512 * 1024 + 1),
        ("notes.txt", 512 * 1024 + 1),
    ],
)
def test_a_source_over_its_read_bound_answers_413_and_writes_nothing(
    client, local_artifact, name, size
):
    folder = local_artifact({name: b"a" * size}, primary=name)
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 413, response.text
    assert _snapshot(folder) == before


def test_a_markdown_source_at_its_read_bound_still_exports(client, local_artifact):
    local_artifact({"report.md": b"a" * (512 * 1024)}, primary="report.md")

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text


@pytest.mark.parametrize("export_format", ["html", "docx"])
def test_markdown_with_a_fence_past_the_cap_answers_413_and_writes_nothing(
    client, local_artifact, export_format
):
    # No primary recorded, so a refusal that came after the record would
    # show in metadata.json.
    fence = "`" * (MAX_FENCE_LENGTH + 1)
    folder = local_artifact({"report.md": f"# Report\n\n{fence}\ncode\n{fence}\n"})
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": export_format})

    assert response.status_code == 413, response.text
    assert response.json()["detail"] == "This file can't be exported."
    assert _snapshot(folder) == before


def test_a_conversion_failure_answers_500_and_writes_nothing(client, local_artifact, monkeypatch):
    folder = local_artifact({"report.md": "# R\n"})
    before = _snapshot(folder)

    def broken(**_kwargs):
        raise ValueError("the writer failed")

    monkeypatch.setattr(artifact_export, "write_docx", broken)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 500, response.text
    assert _snapshot(folder) == before


def _newest(path: Path) -> None:
    """Make *path* the folder's newest file, so the card picks it as primary."""
    later = time.time() + 60
    os.utime(path, (later, later))


def _without_metadata(snapshot: dict[str, object]) -> dict[str, object]:
    return {key: value for key, value in snapshot.items() if key != "metadata.json"}


def test_an_export_with_no_free_name_answers_500_and_writes_no_file(client, local_artifact):
    """The primary is recorded before any name is tried, as the card already
    shows it, and no temporary file stays when no name is free."""
    taken = {"report.docx": b"old", **{f"report ({n}).docx": b"old" for n in range(2, 101)}}
    folder = local_artifact({**taken, "report.md": "# R\n"})
    _newest(folder / "report.md")
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 500, response.text
    assert _without_metadata(_snapshot(folder)) == _without_metadata(before)
    assert _metadata(folder)["primary"] == "report.md"


def test_the_primary_is_recorded_before_the_new_file_takes_its_name(
    client, local_artifact, monkeypatch
):
    """A process stopped between the two steps leaves the card as it was: the
    new report.html never stands beside report.md with no primary recorded."""
    folder = local_artifact({"report.md": "# R\n"})
    place = artifact_export._place_temporary
    recorded: list[object] = []

    def spy(**kwargs):
        recorded.append(_metadata(folder).get("primary"))
        return place(**kwargs)

    monkeypatch.setattr(artifact_export, "_place_temporary", spy)

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text
    assert recorded == ["report.md"]


def test_a_temporary_name_that_stays_takes_the_placed_file_with_it(
    client, local_artifact, monkeypatch
):
    """The placed file shares its inode with the temporary name until that
    goes, and the drafts route refuses a file with two links. So when the
    temporary name can't be removed, no file is left at an output name."""
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    before = _snapshot(folder)
    unlink = artifact_export.dir_unlink

    def stuck(parent, name):
        if name.startswith(".cowork-export-"):
            raise OSError(errno.EIO, "Input/output error")
        unlink(parent, name)

    monkeypatch.setattr(artifact_export, "dir_unlink", stuck)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 500, response.text
    after = _snapshot(folder)
    (left,) = set(after) - set(before)
    assert left.startswith(".cowork-export-")
    assert {key: after[key] for key in before} == before
    assert (folder / left).stat().st_nlink == 1


def test_a_failure_to_record_the_primary_writes_nothing(client, local_artifact, monkeypatch):
    folder = local_artifact({"report.md": "# R\n"})
    before = _snapshot(folder)

    def broken(**_kwargs):
        raise OSError("metadata.json can't be written")

    monkeypatch.setattr(artifact_export, "_record_primary", broken)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 500, response.text
    assert _snapshot(folder) == before


def test_an_html_source_past_the_markdown_bound_still_exports(client, local_artifact):
    page = "<html><body>" + "<p>row</p>" * (2 * MIB // 10 + 1000) + "</body></html>"
    assert 2 * MIB < len(page) < 10 * MIB
    local_artifact({"report.html": page}, primary="report.html")

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 200, response.text


def test_a_symlinked_primary_is_never_read(client, local_artifact, tmp_path):
    outside = tmp_path / "outside.md"
    outside.write_text("# SECRET-OUTSIDE\n", encoding="utf-8")
    folder = local_artifact({"report.md": "# R\n"}, primary="report.md")
    (folder / "report.md").unlink()
    (folder / "report.md").symlink_to(outside)
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code in (404, 422), response.text
    assert _snapshot(folder) == before


def _published(folder: Path, primary: str) -> None:
    (folder / ".published.json").write_text(
        json.dumps(
            {
                primary: {
                    "report_id": "report-1",
                    "url": "https://view.example/r/1",
                    "published": True,
                    "mode": "restricted",
                    "emails": ["a@example.com"],
                    "published_mtime": time.time() + 60,
                }
            }
        ),
        encoding="utf-8",
    )


def _card(folder: Path) -> dict:
    source = ProjectArtifacts(
        base=folder.parent,
        project_id=None,
        project_name="project",
        trusted_anchor=folder.parent.parent.parent,
        root_parts=(".anton", "artifacts"),
    )
    (card,) = artifacts_service.list_artifacts([source])
    # A card with no primary file has no draftUrl.
    return {key: card.get(key) for key in ("path", "publishedUrl", "accessMode", "draftUrl")}


@pytest.mark.parametrize("recorded", [True, False], ids=["primary-recorded", "no-primary"])
@pytest.mark.parametrize(
    ("primary", "files", "exports"),
    [
        ("report.md", {"report.md": "# R\n"}, ["docx", "html", "docx"]),
        ("index.html", {"index.html": "<h1>P</h1>", "notes.md": "# N\n"}, ["docx", "docx"]),
    ],
    ids=["markdown", "html"],
)
def test_exports_leave_the_card_primary_and_publish_state_alone(
    client, local_artifact, primary, files, exports, recorded
):
    # With no recorded primary, the card shows the only HTML file, or else the
    # newest file, so each new export could take the card over.
    folder = local_artifact(files, primary=primary if recorded else None)
    _published(folder, primary)
    card_before = _card(folder)
    metadata_mtime = (folder / "metadata.json").stat().st_mtime_ns
    assert card_before["path"] == str(folder / primary)

    for export_format in exports:
        response = client.post(EXPORT_URL, json={"format": export_format})
        assert response.status_code == 200, response.text
        time.sleep(0.01)

    assert _card(folder) == card_before
    assert _metadata(folder)["primary"] == primary
    assert (folder / "metadata.json").stat().st_mtime_ns == metadata_mtime


def test_a_recorded_primary_is_never_changed(client, local_artifact):
    """A primary that names a file stays as it is, even where the card's own
    fallback would pick another file, and the export writes no metadata."""
    folder = local_artifact({"notes.md": "# N\n", "index.html": "<h1>I</h1>"}, primary="notes.md")
    metadata_before = (folder / "metadata.json").read_bytes()

    response = client.post(EXPORT_URL, json={"format": "html"})

    assert response.status_code == 200, response.text
    assert response.json()["filename"] == "notes.html"
    assert (folder / "metadata.json").read_bytes() == metadata_before


@pytest.mark.parametrize(("export_format", "extension"), [("docx", ".docx"), ("html", ".html")])
def test_a_primary_that_names_no_file_counts_as_unset(
    client, local_artifact, export_format, extension
):
    """The agent renamed report.md to final.md, and metadata.json still names
    report.md. The card shows final.md, and the export records it, so the new
    file never takes the card over and Export stays offered."""
    folder = local_artifact({"final.md": "# F\n"}, primary="report.md")
    _published(folder, "final.md")
    card_before = _card(folder)
    assert card_before["path"] == str(folder / "final.md")

    first = client.post(EXPORT_URL, json={"format": export_format})
    time.sleep(0.01)
    second = client.post(EXPORT_URL, json={"format": export_format})

    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert [first.json()["filename"], second.json()["filename"]] == [
        f"final{extension}",
        f"final (2){extension}",
    ]
    assert _card(folder) == card_before
    assert _metadata(folder)["primary"] == "final.md"


@pytest.mark.parametrize(
    ("files", "status", "card_path"),
    [({"report.md": "# R\n"}, 200, "report.md"), ({}, 422, "")],
    ids=["another-file", "no-other-file"],
)
def test_a_primary_that_names_a_symlink_loop_falls_back_like_the_card(
    client, local_artifact, files, status, card_path
):
    """Python 3.12 raises RuntimeError resolving a link to itself. The card
    and the export both fall back past such a primary instead."""
    folder = local_artifact(files, primary="loop")
    (folder / "loop").symlink_to("loop")

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == status, response.text
    assert _card(folder)["path"] == (str(folder / card_path) if card_path else str(folder))
    if status == 200:
        assert response.json()["filename"] == "report.docx"
        assert _metadata(folder)["primary"] == "report.md"


@pytest.mark.parametrize(
    "hint", ["//203.0.113.5/share/x.md", "\\\\203.0.113.5\\share\\x.md"], ids=["slashes", "backslashes"]
)
def test_a_unc_primary_is_passed_over_before_any_file_call(client, local_artifact, monkeypatch, hint):
    """On Windows, a UNC hint replaces the folder when joined onto it, and
    resolving it opens a connection to that host with the user's credentials.
    The export falls back to the card's own pick without resolving it."""
    folder = local_artifact({"report.md": "# R\n"}, primary=hint)
    resolved: list[str] = []
    resolve = Path.resolve

    def spy(self, *args, **kwargs):
        resolved.append(str(self))
        return resolve(self, *args, **kwargs)

    monkeypatch.setattr(Path, "resolve", spy)

    response = client.post(EXPORT_URL, json={"format": "docx"})

    assert response.status_code == 200, response.text
    assert response.json()["filename"] == "report.docx"
    assert resolved, "the export should resolve its own paths through the spy"
    assert [path for path in resolved if "203.0.113.5" in path] == []
    assert _metadata(folder)["primary"] == "report.md"


def test_permission_is_checked_before_any_file_is_resolved(monkeypatch):
    from cowork.db.scoped import get_scoped_session
    from cowork.services import product_permissions

    async def denied(_scope, _permission):
        return False

    monkeypatch.setattr(product_permissions, "has_product_permission", denied)
    resolve = Mock(side_effect=AssertionError("resolved an artifact"))
    monkeypatch.setattr(workspace_ep, "_owner_workspace", resolve)
    app = FastAPI()
    app.include_router(workspace_ep.router)
    org = TenantScope(org_mode=True, org_id="org", user_id="user")
    app.dependency_overrides[get_scoped_session] = lambda: SimpleNamespace(scope=org)

    with TestClient(app) as test_client:
        response = test_client.post(
            f"/workspace/project/{ARTIFACT_ID}/export", json={"format": "docx"}
        )

    assert response.status_code == 403, response.text
    resolve.assert_not_called()


# ── load ─────────────────────────────────────────────────────────────────────


@contextmanager
def _held_export(monkeypatch):
    """Make every export block in its worker thread until the test lets it go."""
    entered = threading.Event()
    release = threading.Event()

    def held(**kwargs):
        entered.set()
        assert release.wait(timeout=10)
        return ExportResult(
            filename="report.docx",
            draft_url="/api/v1/artifacts/drafts/local/x/report.docx",
            serve_url="",
            omitted_count=0,
            scripts_dropped=False,
        )

    monkeypatch.setattr(workspace_ep, "export_primary_file", held)
    try:
        yield SimpleNamespace(entered=entered, release=release)
    finally:
        release.set()


@pytest.mark.parametrize("second_format", ["docx", "html"])
def test_a_second_export_while_one_runs_answers_429_at_once(
    client, local_artifact, monkeypatch, second_format
):
    local_artifact({"report.md": "# R\n"}, primary="report.md")
    answers: list[int] = []
    with _held_export(monkeypatch) as held:
        first = threading.Thread(
            target=lambda: answers.append(
                client.post(EXPORT_URL, json={"format": "docx"}).status_code
            )
        )
        first.start()
        assert held.entered.wait(timeout=10)

        started = time.monotonic()
        second = client.post(EXPORT_URL, json={"format": second_format})
        elapsed = time.monotonic() - started

        held.release.set()
        first.join(timeout=10)

    assert second.status_code == 429, second.text
    assert second.json()["detail"] == "An export is already running; try again in a moment."
    assert elapsed < 1.0
    assert answers == [200]
    # The slot is free again once the first export answers.
    after = client.post(EXPORT_URL, json={"format": "html"})
    assert after.status_code == 200, after.text


# ── org tenancy ──────────────────────────────────────────────────────────────


@contextmanager
def _scoped(org_id: str, user_id: str):
    with Session(get_engine(get_app_settings().database.uri)) as raw:
        yield ScopedSession(raw, TenantScope(org_mode=True, org_id=org_id, user_id=user_id))


def _write_artifact(base: Path, slug: str, files: dict[str, str | bytes], **metadata) -> tuple[str, Path]:
    artifact_id = uuid4().hex
    folder = base / slug
    for name, content in files.items():
        path = folder / name
        path.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(content, bytes):
            path.write_bytes(content)
        else:
            path.write_text(content, encoding="utf-8")
    (folder / "metadata.json").write_text(
        json.dumps({"id": artifact_id, "slug": slug, "name": slug, **metadata}), encoding="utf-8"
    )
    return artifact_id, folder


@pytest.fixture
def org_world(tmp_path, org_deployment):
    """One org project: the creator's report, a member's report, an ownerless
    page, and a report in the creator's private conversation workspace."""
    org_id, creator, member = str(uuid4()), str(uuid4()), str(uuid4())
    project = make_project(tmp_path, org_id)
    shared = project_root(project)
    private = legacy_root(project, str(make_conversation(project, creator)))
    artifacts = {}
    for slug, owner, files, primary, root in (
        ("mine", creator, {"report.md": "# Mine\n"}, "report.md", shared),
        ("theirs", member, {"report.md": "# Theirs\n"}, "report.md", shared),
        ("page", creator, {"index.html": "<h1>Page</h1>"}, "index.html", shared),
        ("orphan", None, {"report.md": "# Orphan\n"}, "report.md", shared),
        ("private", creator, {"report.md": "# Private\n"}, "report.md", private),
    ):
        artifact_id, folder = _write_artifact(root.base, slug, files, type="document", primary=primary)
        artifacts[slug] = SimpleNamespace(id=artifact_id, folder=folder)
        if owner is not None and root is shared:
            with _scoped(org_id, owner) as session:
                ownership.record_artifact_owner(session, project.id, slug, owner, action="create")
    return SimpleNamespace(
        org_id=org_id, creator=creator, member=member, project=project, artifacts=artifacts
    )


@pytest.fixture
def role(monkeypatch):
    """Answer the live ``artifact.manage`` check as each test says."""
    from cowork.services import product_permissions

    decision = SimpleNamespace(artifact_manage=True)

    async def decide(_scope, permission):
        return decision.artifact_manage if permission == "artifact.manage" else True

    monkeypatch.setattr(product_permissions, "has_product_permission", decide)
    monkeypatch.setattr(workspace_ep, "has_product_permission", decide)
    return decision


@pytest.fixture
def org_export_on(monkeypatch):
    """Turn org export on, as the release that lifts the switch will."""
    monkeypatch.setattr(artifact_export, "ORG_EXPORT_ENABLED", True)


async def _export(world, user_id: str, slug: str, export_format: str = "docx", *, org_id=None):
    with _scoped(org_id or world.org_id, user_id) as session:
        return await workspace_ep.export_workspace_artifact(
            str(world.project.id),
            world.artifacts[slug].id,
            ExportRequest(format=export_format),
            session,
        )


async def test_the_owner_with_artifact_manage_exports_and_gets_no_serve_url(
    org_world, role, org_export_on
):
    result = await _export(org_world, org_world.creator, "mine")

    folder = org_world.artifacts["mine"].folder
    assert result.model_dump(by_alias=True) == {
        "filename": "report.docx",
        "draftUrl": (
            f"/api/v1/artifacts/drafts/{org_world.project.id}/"
            f"{org_world.artifacts['mine'].id}/report.docx"
        ),
        "serveUrl": "",
        "omittedCount": 0,
        "scriptsDropped": False,
    }
    docx.Document(io.BytesIO((folder / "report.docx").read_bytes()))


@pytest.mark.parametrize(
    ("who", "slug", "export_format", "status"),
    [
        ("member", "mine", "docx", 403),
        # Permission comes before format: the page offers no HTML export.
        ("member", "page", "html", 403),
        ("creator", "orphan", "docx", 403),
        ("member", "private", "docx", 404),
        ("stranger", "mine", "docx", 404),
    ],
    ids=["reviewer", "reviewer-asking-a-format-not-offered", "ownerless", "no-grant", "other-org"],
)
@pytest.mark.parametrize("org_export", [False, True], ids=["org-export-off", "org-export-on"])
async def test_only_the_owner_may_export(
    org_world, role, monkeypatch, org_export, who, slug, export_format, status
):
    """Each refusal comes before the org export switch, so it holds either way."""
    monkeypatch.setattr(artifact_export, "ORG_EXPORT_ENABLED", org_export)
    folder = org_world.artifacts[slug].folder
    before = _snapshot(folder)
    user_id = {"member": org_world.member, "creator": org_world.creator, "stranger": str(uuid4())}[who]
    org_id = str(uuid4()) if who == "stranger" else None

    with pytest.raises(HTTPException) as refused:
        await _export(org_world, user_id, slug, export_format, org_id=org_id)

    assert refused.value.status_code == status
    assert _snapshot(folder) == before


async def test_an_owner_without_artifact_manage_gets_403(org_world, role):
    role.artifact_manage = False
    folder = org_world.artifacts["mine"].folder
    before = _snapshot(folder)

    with pytest.raises(HTTPException) as refused:
        await _export(org_world, org_world.creator, "mine")

    assert refused.value.status_code == 403
    assert _snapshot(folder) == before


@pytest.mark.parametrize(("slug", "export_format"), [("mine", "docx"), ("page", "html")])
async def test_while_org_export_is_off_the_owner_gets_403_and_nothing_is_written(
    org_world, role, slug, export_format
):
    folder = org_world.artifacts[slug].folder
    before = _snapshot(folder)

    with pytest.raises(HTTPException) as refused:
        await _export(org_world, org_world.creator, slug, export_format)

    assert refused.value.status_code == 403
    assert _snapshot(folder) == before


async def test_the_owner_asking_a_format_not_offered_gets_422(org_world, role, org_export_on):
    folder = org_world.artifacts["page"].folder
    before = _snapshot(folder)

    with pytest.raises(HTTPException) as refused:
        await _export(org_world, org_world.creator, "page", "html")

    assert refused.value.status_code == 422
    assert _snapshot(folder) == before


@pytest.mark.parametrize(("org_mode", "budget"), [(True, 10 * MIB), (False, 25 * MIB)])
async def test_the_word_inline_budget_is_smaller_in_org_mode(
    org_world, role, org_export_on, monkeypatch, org_mode, budget
):
    if not org_mode:
        monkeypatch.setenv("COWORK_TENANCY_MODE", "local")
        get_app_settings.cache_clear()
    budgets: list[int] = []
    inline = artifact_export.inline_artifact_html

    def spy(**kwargs):
        budgets.append(kwargs["limits"].max_total_bytes)
        return inline(**kwargs)

    monkeypatch.setattr(artifact_export, "inline_artifact_html", spy)
    folder = org_world.artifacts["mine"].folder
    source = ProjectArtifacts(
        base=folder.parent,
        project_id=str(org_world.project.id),
        project_name=org_world.project.name,
        trusted_anchor=Path(org_world.project.path),
        root_parts=(".anton", "artifacts"),
    )

    artifact_export.export_primary_file(
        source=source,
        folder_name=folder.name,
        metadata=_metadata(folder),
        artifact_id=org_world.artifacts["mine"].id,
        export_format="docx",
        can_edit=True,
    )

    assert budgets == [budget]


def padded_png(size: int, *, filler: bytes = b"x") -> bytes:
    """A valid 1x1 PNG of *size* bytes, padded with a private ancillary chunk
    that repeats *filler*."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    head = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
    tail = chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00")) + chunk(b"IEND", b"")
    padding = size - len(head) - len(tail) - 12
    filled = (filler * (padding // len(filler) + 1))[:padding]
    return head + chunk(b"teXt", filled) + tail


async def test_org_assets_past_the_budget_are_dropped_and_counted(org_world, role, org_export_on):
    folder = org_world.artifacts["mine"].folder
    # Each image is under the 2 MiB asset cap even as base64; together they
    # are over the 10 MiB budget.
    images = {f"chart{i}.png": padded_png(int(1.4 * MIB)) for i in range(8)}
    for name, data in images.items():
        (folder / name).write_bytes(data)
    markdown = "# Charts\n\n" + "\n\n".join(f"![chart]({name})" for name in images) + "\n"
    (folder / "report.md").write_text(markdown, encoding="utf-8")

    result = await _export(org_world, org_world.creator, "mine")

    assert result.omitted_count >= 1
    document = docx.Document(io.BytesIO((folder / result.filename).read_bytes()))
    assert 1 <= len(document.inline_shapes) <= 7


async def test_one_export_per_user_and_org_at_a_time(org_world, role, org_export_on, monkeypatch):
    """While the creator's export runs, the creator's second export and a
    co-member's export of their own artifact both answer 429 at once."""
    import asyncio

    with _held_export(monkeypatch) as held:
        first = asyncio.create_task(_export(org_world, org_world.creator, "mine"))
        assert await asyncio.to_thread(held.entered.wait, 10)

        for user_id, slug in ((org_world.creator, "mine"), (org_world.member, "theirs")):
            started = time.monotonic()
            with pytest.raises(HTTPException) as busy:
                await _export(org_world, user_id, slug)
            assert busy.value.status_code == 429
            assert time.monotonic() - started < 1.0

        held.release.set()
        result = await first

    assert result.filename == "report.docx"


# ── PDF through a render service ─────────────────────────────────────────────
#
# The render service is an ``httpx.MockTransport`` handed to the export's
# client, so these tests see the exact request the route sends and answer it
# the way a render service can: with a PDF, busy, refusing, redirecting or
# broken.

PDF_RESULT = json.loads((Path(__file__).parent / "fixtures" / "render_pdf" / "response.json").read_text(encoding="utf-8"))
PDF_FAILURES = yaml.safe_load(
    (Path(__file__).parent / "fixtures" / "artifact_export" / "pdf_failures.yaml").read_text(encoding="utf-8")
)["cases"]
PDF_SETTINGS = ExportPdfSettings(engine="render", render_url="http://render.test")

def _streamed(*, status: int, body: bytes) -> httpx.Response:
    """An answer the client reads as a stream, as it reads a real one."""
    return httpx.Response(
        status, stream=httpx.ByteStream(body), headers={"Content-Type": "application/json"}
    )


class RenderService:
    """Answers each request with the next queued answer and keeps every request."""

    def __init__(self) -> None:
        self.requests: list[RenderPdfRequest] = []
        self.answers: list[dict] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(RenderPdfRequest.model_validate_json(request.content))
        answer = self.answers.pop(0) if self.answers else {"status": 200}
        status = answer["status"]
        if status == 302:
            return httpx.Response(302, headers={"Location": answer["location"]})
        if "raw" in answer:
            return _streamed(status=status, body=answer["raw"].encode())
        if status == 200:
            result = {**PDF_RESULT, **({"pdf_base64": answer["pdf_base64"]} if "pdf_base64" in answer else {})}
            return _streamed(status=200, body=json.dumps(result).encode())
        return httpx.Response(status)


@pytest.fixture
def render_service(monkeypatch):
    """Turn PDF export on, with this stand-in as the render service."""
    service = RenderService()
    monkeypatch.setattr(artifact_export, "export_pdf_settings", lambda: PDF_SETTINGS)
    monkeypatch.setattr(
        artifact_export,
        "_pdf_client",
        lambda: RenderServicePdf(settings=PDF_SETTINGS, transport=httpx.MockTransport(service)),
    )
    return service


def test_exports_read_the_pdf_settings_from_the_process_environment(monkeypatch):
    monkeypatch.setenv("COWORK_EXPORT_PDF_ENGINE", "render")
    monkeypatch.setenv("COWORK_EXPORT_RENDER_URL", "http://render.test")
    get_export_pdf_settings.cache_clear()
    try:
        assert artifact_export.export_pdf_settings() == PDF_SETTINGS
    finally:
        get_export_pdf_settings.cache_clear()


def test_pdf_is_offered_only_with_a_render_service_and_never_in_org_mode(monkeypatch):
    report = ExportSource(source_name="report.md", is_backend=False)
    page = ExportSource(source_name="page.html", is_backend=False)
    sheet = ExportSource(source_name="data.csv", is_backend=False)

    def offered(source: ExportSource, *, org_mode: bool = False) -> list[str]:
        return artifact_export.offered_export_formats(source=source, can_edit=True, org_mode=org_mode)

    assert offered(report) == ["html", "docx"]
    monkeypatch.setattr(artifact_export, "export_pdf_settings", lambda: PDF_SETTINGS)
    assert offered(report) == ["html", "docx", "pdf"]
    assert offered(page) == ["docx", "pdf"]
    assert offered(sheet) == []
    assert offered(report, org_mode=True) == []


def test_a_pdf_export_writes_the_render_services_pdf(client, local_artifact, render_service):
    folder = local_artifact({"report.md": "# Report\n\nRevenue rose.\n"})

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 200, response.text
    assert response.json()["filename"] == "report.pdf"
    written = (folder / "report.pdf").read_bytes()
    assert written.startswith(b"%PDF-")
    [sent] = render_service.requests
    assert sent.scripts is False
    assert "Revenue rose." in sent.html


@pytest.mark.parametrize(("primary", "scripts"), [("page.html", True), ("page.htm", True), ("notes.txt", False)])
def test_only_an_html_page_asks_for_its_scripts_to_run(
    client, local_artifact, render_service, primary, scripts
):
    local_artifact({primary: "<p>hello</p>" if primary.startswith("page") else "hello"}, primary=primary)

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 200, response.text
    assert [request.scripts for request in render_service.requests] == [scripts]


def test_a_pdf_inlines_at_most_its_own_smaller_budget(client, local_artifact, render_service):
    images = {f"chart{n}.png": padded_png(int(1.9 * MIB), filler=bytes([65 + n])) for n in range(4)}
    page = "".join(f'<img src="chart{n}.png">' for n in range(4))
    local_artifact({**images, "report.html": f"<html><body>{page}</body></html>"}, primary="report.html")

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 200, response.text
    assert response.json()["omittedCount"] >= 1
    [sent] = render_service.requests
    assert len(sent.html) < 1.4 * artifact_export.PDF_INLINE_BUDGET_BYTES


def test_a_pdf_of_an_html_page_keeps_its_scripts(client, local_artifact, render_service):
    cdn = "https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js"
    page = f'<html><body><p>chart</p><script>window.x = 1;</script><script src="{cdn}"></script></body></html>'
    local_artifact({"page.html": page}, primary="page.html")

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 200, response.text
    [sent] = render_service.requests
    assert "window.x = 1;" in sent.html
    assert f'src="{cdn}"' in sent.html
    assert response.json()["scriptsDropped"] is False


@pytest.mark.parametrize("case", PDF_FAILURES, ids=[case["id"] for case in PDF_FAILURES])
def test_a_pdf_that_does_not_come_back_writes_nothing(client, local_artifact, render_service, case):
    folder = local_artifact({"report.md": "# Report\n"})
    before = _snapshot(folder)
    render_service.answers = [dict(answer) for answer in case["answers"]]

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == case["status"], response.text
    assert response.json()["detail"] == case["detail"]
    assert _snapshot(folder) == before


def test_no_pdf_comes_back_in_time_answers_504_and_writes_nothing(
    client, local_artifact, monkeypatch
):
    settings = PDF_SETTINGS.model_copy(update={"timeout_seconds": 0.5})

    async def stall(_request: httpx.Request) -> httpx.Response:
        await anyio.sleep(5)
        return httpx.Response(503)

    monkeypatch.setattr(artifact_export, "export_pdf_settings", lambda: settings)
    monkeypatch.setattr(
        artifact_export,
        "_pdf_client",
        lambda: RenderServicePdf(settings=settings, transport=httpx.MockTransport(stall)),
    )
    folder = local_artifact({"report.md": "# Report\n"})
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 504, response.text
    assert _snapshot(folder) == before


def test_without_a_render_service_pdf_answers_422_and_nothing_makes_one(client, local_artifact):
    folder = local_artifact({"report.md": "# Report\n"})
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 422, response.text
    assert _snapshot(folder) == before


def test_a_pdf_of_markdown_with_a_fence_past_the_cap_answers_413_and_sends_nothing(
    client, local_artifact, render_service
):
    fence = "`" * (MAX_FENCE_LENGTH + 1)
    folder = local_artifact({"report.md": f"# Report\n\n{fence}\ncode\n{fence}\n"})
    before = _snapshot(folder)

    response = client.post(EXPORT_URL, json={"format": "pdf"})

    assert response.status_code == 413, response.text
    assert response.json()["detail"] == "This file can't be exported."
    assert render_service.requests == []
    assert _snapshot(folder) == before


def _local_person(user_id: str) -> TenantScope:
    return TenantScope(org_mode=False, org_id=None, user_id=user_id)


async def test_pdf_export_has_its_own_slot():
    """A running PDF export leaves Word and HTML export free, and only one PDF runs."""
    async with export_slot(scope=_local_person("alice"), export_format="pdf"):
        async with export_slot(scope=_local_person("bob"), export_format="docx"):
            pass
        async with export_slot(scope=_local_person("dave"), export_format="html"):
            pass
        with pytest.raises(ExportRefused):
            async with export_slot(scope=_local_person("carol"), export_format="pdf"):
                pass
    async with export_slot(scope=_local_person("carol"), export_format="pdf"):
        await anyio.sleep(0)
