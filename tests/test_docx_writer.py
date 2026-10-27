"""The Word writer: structure, links, images, limits, and no fetch or file read."""
import ast
import base64
import importlib.machinery
import io
import os
import struct
import sys
import threading
import types
import zipfile
import zlib
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import docx
import pytest
import yaml
from docx.document import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from cowork.common.paths import pinned_dir
from cowork.schemas.artifact_export import InlinedDocument, InlineLimits, InlineMode
from cowork.services import artifact_inline, docx_writer
from cowork.services.artifact_inline import inline_artifact_html
from cowork.services.docx_writer import (
    DEADLINE_SECONDS,
    MAX_DEPTH,
    MAX_NODES,
    MAX_TABLE_CELLS,
    MAX_TABLE_COLUMNS,
    TABLE_CELLS_NOTE,
    TRUNCATION_NOTE,
    write_docx,
)
from tests._no_fetch import no_fetch, recording_server

FIXTURES = Path(__file__).parent / "fixtures" / "docx_writer"
SENTINEL = b"SENTINEL-DOCX-WRITER-5d1c"
FORBIDDEN_MARKUP = (b"fldSimple", b"fldChar", b"instrText", b"altChunk")
_RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_SAFE_XML = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
# More digits than Python converts to an int; fixtures write it as {digits}.
_HUGE_NUMBER = "9" * 5000
# docx.Document() reads python-docx's bundled template. Besides it, only the
# import system may open files: modules a library imports on first use.
_DOCX_TEMPLATES = Path(docx.__file__).resolve().parent / "templates"
_MODULE_SUFFIXES = frozenset({*importlib.machinery.all_suffixes(), ".so", ".pyd"})
_IMPORT_ROOTS = tuple(Path(entry).resolve() for entry in sys.path if entry)


class _OpenRecorder:
    """Collects the path of every file opened while recording.

    It records the recording thread and any thread started while recording,
    so a writer that hands its I/O to a thread is still caught. Threads that
    were already running belong to other tests, such as a service one of
    them left writing in the background, and are ignored.

    An audit hook can't be removed, so one recorder stays installed for the
    session and records only inside ``recording()``.
    """

    def __init__(self) -> None:
        self._paths: list[str] | None = None
        self._ignored: frozenset[threading.Thread] = frozenset()

    def __call__(self, event: str, args: tuple) -> None:
        # Thread objects, not idents: a thread started while recording can
        # reuse the ident of one that was running before and has ended.
        if event == "open" and self._paths is not None and threading.current_thread() not in self._ignored:
            path = args[0]
            if isinstance(path, int):
                # A descriptor names no file; the open that made it named one.
                return
            self._paths.append(os.fsdecode(path) if isinstance(path, str | bytes | os.PathLike) else repr(path))

    @contextmanager
    def recording(self) -> Iterator[list[str]]:
        self._ignored = frozenset(threading.enumerate()) - {threading.current_thread()}
        self._paths = paths = []
        try:
            yield paths
        finally:
            self._paths = None
            self._ignored = frozenset()


_OPENS = _OpenRecorder()
sys.addaudithook(_OPENS)


def test_the_open_recorder_ignores_threads_that_were_already_running(tmp_path):
    """A thread another test left running may open a file at any moment; a
    thread started while recording is still caught."""
    go, done = threading.Event(), threading.Event()

    def background() -> None:
        go.wait(5)
        (tmp_path / "background.txt").write_text("x")
        done.set()

    earlier = threading.Thread(target=background)
    earlier.start()
    with _OPENS.recording() as opened:
        go.set()
        assert done.wait(5)
        started = threading.Thread(target=lambda: (tmp_path / "started.txt").write_text("y"))
        started.start()
        started.join()
    earlier.join()

    names = {Path(path).name for path in opened}
    assert names & {"background.txt", "started.txt"} == {"started.txt"}


def _opened_outside_the_template(*, paths: list[str]) -> list[str]:
    """The opened paths that are neither python-docx's template nor an imported module."""
    unexpected: list[str] = []
    for path in paths:
        resolved = Path(path).resolve()
        if resolved.is_relative_to(_DOCX_TEMPLATES):
            continue
        in_import_roots = any(resolved.is_relative_to(root) for root in _IMPORT_ROOTS)
        if resolved.suffix in _MODULE_SUFFIXES and in_import_roots:
            continue
        # The import system writes a module's bytecode cache on its first import.
        if "__pycache__" in resolved.parts and in_import_roots:
            continue
        unexpected.append(path)
    return unexpected


def _cases(name: str) -> list[dict]:
    return yaml.safe_load((FIXTURES / name).read_text(encoding="utf-8"))


def _ids(cases: list[dict]) -> list[str]:
    return [case["name"] for case in cases]


@dataclass(frozen=True)
class Package:
    document: WordDocument
    parts: dict[str, bytes]

    @property
    def text(self) -> str:
        """Every paragraph's text, table cells included, one per line."""
        return "\n".join(self.paragraph_texts())

    def paragraph_texts(self) -> list[str]:
        """Every paragraph's text in document order, table cells included."""
        body = self.document.element.body
        return [Paragraph(p, self.document._body).text for p in body.iter(qn("w:p"))]

    def media(self) -> dict[str, bytes]:
        return {name: body for name, body in self.parts.items() if name.startswith("word/media/")}

    def external_relationships(self) -> list[etree._Element]:
        return [
            relationship
            for name, body in self.parts.items()
            if name.endswith(".rels")
            for relationship in etree.fromstring(body, _SAFE_XML)
            if relationship.get("TargetMode") == "External"
        ]


def _inlined(html: str) -> InlinedDocument:
    return InlinedDocument(html=html, omitted=[], inlined_bytes=0, scripts_dropped=False)


def _open_package(data: bytes) -> Package:
    """Reload the bytes and check the rules every package keeps."""
    document = docx.Document(io.BytesIO(data))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, body in parts.items():
        if name.endswith((".xml", ".rels")):
            etree.fromstring(body, _SAFE_XML)
            for markup in FORBIDDEN_MARKUP:
                assert markup not in body, f"{markup!r} in {name}"
    package = Package(document=document, parts=parts)
    for relationship in package.external_relationships():
        assert relationship.get("Type") == RT.HYPERLINK
    for tc in document.element.body.iter(qn("w:tc")):
        assert tc[-1].tag == qn("w:p"), "a table cell must end with a paragraph"
    return package


def _write(html: str, *, title: str = "Fallback title") -> Package:
    with no_fetch() as recorder, _OPENS.recording() as opened:
        data = write_docx(document=_inlined(html), title=title)
    assert recorder.calls == []
    assert _opened_outside_the_template(paths=opened) == []
    return _open_package(data)


def _png(*, width: int, height: int, text: bytes = b"") -> bytes:
    """A grayscale PNG with no pHYs chunk, built without an imaging library.

    ``text`` goes in a tEXt chunk, which pads the file without adding pixels.
    """

    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))

    scanlines = (b"\x00" + b"\x80" * width) * height
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0))
        + (chunk(b"tEXt", b"Comment\x00" + text) if text else b"")
        + chunk(b"IDAT", zlib.compress(scanlines))
        + chunk(b"IEND", b"")
    )


def _png_data_uri(*, width: int, height: int) -> str:
    return "data:image/png;base64," + base64.b64encode(_png(width=width, height=height)).decode()


def _padded_png(*, index: int, size: int) -> bytes:
    """A 1x1 PNG of about ``size`` bytes, distinct for each ``index``."""
    return _png(width=1, height=1, text=f"{index:02d}".encode() * (size // 2))


# The criteria fixture


@pytest.fixture(scope="module")
def structure() -> tuple[Package, dict]:
    html = (FIXTURES / "structure.html").read_text(encoding="utf-8")
    return _write(html, title="Fallback title"), _cases("structure.yaml")


def test_headings_map_to_heading_styles(structure):
    package, expected = structure
    headings = [
        {"level": int(p.style.name.removeprefix("Heading ")), "text": p.text}
        for p in package.document.paragraphs
        if p.style.name.startswith("Heading ")
    ]
    assert headings == expected["headings"]


def test_head_title_becomes_the_core_title(structure):
    package, expected = structure
    assert package.document.core_properties.title == expected["title"]


def test_pre_keeps_its_whitespace_in_a_monospace_run(structure):
    package, expected = structure
    pre = next(p for p in package.document.paragraphs if p.text == expected["pre"])
    assert [run.font.name for run in pre.runs] == [expected["monospace_font"]]
    assert pre.runs[0].text == expected["pre"]


def test_lists_use_list_styles_levels_and_one_numbering_per_list(structure):
    package, expected = structure
    document = package.document
    items = [p for p in document.paragraphs if p.style.name.startswith(("List Bullet", "List Number"))]
    assert [{"text": p.text.strip(), "style": p.style.name} for p in items] == [
        {"text": item["text"], "style": item["style"]} for item in expected["list_items"]
    ]

    numbering = document.part.numbering_part.element
    num_ids_by_list: dict[str, set[int]] = {}
    for paragraph, item in zip(items, expected["list_items"], strict=True):
        num_pr = paragraph._p.pPr.numPr
        assert num_pr.ilvl.val == 0
        num_ids_by_list.setdefault(item["list"], set()).add(num_pr.numId.val)
        style_num_id = document.styles[item["style"]].element.pPr.numPr.numId.val
        clone = numbering.num_having_numId(num_pr.numId.val)
        assert num_pr.numId.val != style_num_id
        assert clone.abstractNumId.val == numbering.num_having_numId(style_num_id).abstractNumId.val
        starts = [override.startOverride.val for override in clone.lvlOverride_lst]
        assert starts == ([] if item["start"] is None else [item["start"]])

    assert all(len(num_ids) == 1 for num_ids in num_ids_by_list.values())
    assert len({num_id for num_ids in num_ids_by_list.values() for num_id in num_ids}) == len(num_ids_by_list)


def test_table_uses_table_grid_with_merged_cells(structure):
    package, expected = structure
    table_expected = expected["table"]
    [table] = package.document.tables
    assert table.style.name == table_expected["style"]
    assert (len(table.rows), len(table.columns)) == (table_expected["rows"], table_expected["columns"])
    assert [[cell.text for cell in row.cells] for row in table.rows] == table_expected["cells"]
    for merge in table_expected["merged"]:
        assert table.cell(*merge["from"])._tc is table.cell(*merge["to"])._tc
    assert table.cell(0, 1)._tc is not table.cell(0, 2)._tc
    assert table.cell(1, 0)._tc is not table.cell(1, 1)._tc


def test_one_hyperlink_with_an_external_relationship(structure):
    package, expected = structure
    link = expected["hyperlink"]
    hyperlinks = list(package.document.element.body.iter(qn("w:hyperlink")))
    assert len(hyperlinks) == 1
    relationship = package.document.part.rels[hyperlinks[0].get(qn("r:id"))]
    assert relationship.is_external
    assert relationship.target_ref == link["url"]
    assert "".join(t.text for t in hyperlinks[0].iter(qn("w:t"))) == link["text"]
    assert any(p.text == link["paragraph"] for p in package.document.paragraphs)
    assert [r.get("Target") for r in package.external_relationships()] == [link["url"]]


def test_document_under_the_limits_has_no_closing_note(structure):
    package, _ = structure
    assert TRUNCATION_NOTE not in package.text


def test_package_carries_no_template_metadata_or_thumbnail():
    before = datetime.now(UTC).replace(microsecond=0)
    package = _write("<p>x</p>")
    after = datetime.now(UTC)
    properties = package.document.core_properties
    assert (properties.author, properties.comments, properties.last_modified_by) == ("", "", "")
    assert before <= properties.created <= after
    assert properties.modified == properties.created
    package_relationships = etree.fromstring(package.parts["_rels/.rels"], _SAFE_XML)
    assert RT.THUMBNAIL not in {relationship.get("Type") for relationship in package_relationships}
    assert not [name for name in package.parts if name.startswith("docProps/thumbnail")]


# Titles and text flow


@pytest.mark.parametrize("case", _cases("titles.yaml"), ids=_ids(_cases("titles.yaml")))
def test_core_title_comes_from_the_head_or_the_argument(case):
    package = _write(case["html"], title=case["title_arg"])
    assert package.document.core_properties.title == case["core_title"]


@pytest.mark.parametrize("case", _cases("long_titles.yaml"), ids=_ids(_cases("long_titles.yaml")))
def test_long_titles_are_cut_to_the_core_property_limit(case):
    long_title = case["unit"] * case["repeat"]
    if case["source"] == "head":
        package = _write(f"<title>{long_title}</title><p>x</p>", title="Argument title")
    else:
        package = _write("<p>x</p>", title=long_title)
    core_title = package.document.core_properties.title
    assert len(core_title) == case["core_length"]
    assert long_title.startswith(core_title)


@pytest.mark.parametrize("case", _cases("text_flow.yaml"), ids=_ids(_cases("text_flow.yaml")))
def test_text_and_blocks_flow_into_paragraphs(case):
    package = _write(case["html"])
    paragraphs = package.document.paragraphs
    assert [[p.style.name, p.text] for p in paragraphs] == case["paragraphs"]
    if "numbered" in case:
        assert [p._p.pPr is not None and p._p.pPr.numPr is not None for p in paragraphs] == case["numbered"]


@pytest.mark.parametrize("case", _cases("inline_styles.yaml"), ids=_ids(_cases("inline_styles.yaml")))
def test_inline_tags_become_formatted_runs(case):
    package = _write(f"<p>plain {case['html']} plain</p>")
    [run] = [run for run in package.document.paragraphs[0].runs if run.text == "styled"]
    font = run.font
    actual = {
        "bold": font.bold,
        "italic": font.italic,
        "underline": font.underline,
        "strike": font.strike,
        "subscript": font.subscript,
        "superscript": font.superscript,
        "font_name": font.name,
        "highlight": font.highlight_color.name if font.highlight_color is not None else None,
        "size_pt": font.size.pt if font.size is not None else None,
    }
    # python-docx reads superscript as False on a subscript run; unset and
    # False both mean "not this formatting".
    assert {key: value or None for key, value in actual.items()} == {key: None for key in actual} | case["expect"]


@pytest.mark.parametrize("case", _cases("list_numbering.yaml"), ids=_ids(_cases("list_numbering.yaml")))
def test_ordered_lists_number_their_items_as_a_browser_does(case):
    package = _write(case["html"].replace("{digits}", _HUGE_NUMBER))
    numbering = package.document.part.numbering_part.element
    shown: list[int] = []
    formats: set[str] = set()
    items_before: dict[int, int] = {}
    for paragraph in package.document.paragraphs:
        num_id = paragraph._p.pPr.numPr.numId.val
        num = numbering.num_having_numId(num_id)
        [override] = num.lvlOverride_lst
        # Word counts each item on from its w:num's start override.
        shown.append(override.startOverride.val + items_before.get(num_id, 0))
        items_before[num_id] = items_before.get(num_id, 0) + 1
        levels = override.findall(qn("w:lvl")) or [
            level
            for abstract in numbering.findall(qn("w:abstractNum"))
            if abstract.get(qn("w:abstractNumId")) == str(num.abstractNumId.val)
            for level in abstract.findall(qn("w:lvl"))
            if level.get(qn("w:ilvl")) == "0"
        ]
        formats.add(levels[0].find(qn("w:numFmt")).get(qn("w:val")))
    assert shown == case["numbers"]
    assert formats == {case["format"]}


# Links


@pytest.mark.parametrize("case", _cases("links.yaml"), ids=_ids(_cases("links.yaml")))
def test_only_http_https_and_mailto_become_hyperlinks(case):
    href = case["href"].replace("&", "&amp;").replace('"', "&quot;")
    package = _write(f'<p>Before <a href="{href}">link text</a> after.</p>')
    [paragraph] = package.document.paragraphs
    assert paragraph.text == "Before link text after."
    hyperlinks = list(package.document.element.body.iter(qn("w:hyperlink")))
    targets = [r.get("Target") for r in package.external_relationships()]
    if case["target"] is None:
        assert hyperlinks == []
        assert targets == []
    else:
        assert len(hyperlinks) == 1
        assert targets == [case["target"]]


def test_links_share_one_relationship_per_url_beside_images():
    image = f'<img src="{_png_data_uri(width=3, height=2)}" alt="chart">'
    package = _write(
        '<p><a href="https://example.com/a">first a</a> <a href="https://example.com/b">b</a></p>'
        f"<p>{image}</p>"
        '<p><a href="https://example.com/a">second a</a> <a href="https://example.com/c">c</a></p>'
    )
    rels = package.document.part.rels
    targets = {
        "".join(t.text for t in hyperlink.iter(qn("w:t"))): rels[hyperlink.get(qn("r:id"))].target_ref
        for hyperlink in package.document.element.body.iter(qn("w:hyperlink"))
    }
    assert targets == {
        "first a": "https://example.com/a",
        "b": "https://example.com/b",
        "second a": "https://example.com/a",
        "c": "https://example.com/c",
    }
    assert sorted(r.get("Target") for r in package.external_relationships()) == [
        "https://example.com/a",
        "https://example.com/b",
        "https://example.com/c",
    ]
    [blip] = package.document.element.body.iter(qn("a:blip"))
    assert rels[blip.get(qn("r:embed"))].reltype == RT.IMAGE


# Images


@pytest.mark.parametrize("case", _cases("images.yaml"), ids=_ids(_cases("images.yaml")))
def test_images_pass_as_bytes_or_become_a_placeholder(case):
    package = _write(f'<p>Chart: <img src="{case["src"]}" alt="{case["alt"]}"></p>')
    shapes = package.document.inline_shapes
    if case["picture"]:
        assert len(shapes) == 1
        assert shapes[0]._inline.docPr.get("descr") == case["alt"]
        payload = case["src"].partition(",")[2]
        # The bytes reach the package unchanged: nothing re-encodes them.
        assert list(package.media().values()) == [base64.b64decode("".join(payload.split()))]
        assert "[image" not in package.text
    else:
        assert len(shapes) == 0
        assert package.media() == {}
        placeholder = f"[image: {case['alt']}]" if case["alt"] else "[image]"
        assert package.text == f"Chart: {placeholder}"


def _exif_jpeg(*, entries: int) -> bytes:
    """A 3x2 JPEG whose Exif directory holds *entries* text entries, all naming one 1 KiB string."""
    text = b"x" * 1023 + b"\x00"
    text_offset = 8 + 2 + 12 * entries + 4
    directory = (
        struct.pack(">H", entries)
        + b"".join(
            struct.pack(">HHII", 0x8000 + index, 2, len(text), text_offset) for index in range(entries)
        )
        + struct.pack(">I", 0)
    )
    tiff = b"MM\x00\x2a" + struct.pack(">I", 8) + directory + text
    app1 = b"\xff\xe1" + struct.pack(">H", 8 + len(tiff)) + b"Exif\x00\x00" + tiff
    sof0 = b"\xff\xc0" + struct.pack(">HBHHB", 11, 8, 2, 3, 1) + b"\x01\x11\x00"
    sos = b"\xff\xda" + struct.pack(">HB", 8, 1) + b"\x01\x00\x00\x3f\x00"
    return b"\xff\xd8" + app1 + sof0 + sos + b"\x00\xff\xd9"


_JFIF_APP0 = b"\xff\xe0" + struct.pack(">H", 16) + b"JFIF\x00\x01\x01\x00\x00\x48\x00\x48\x00\x00"
_SOF0 = b"\xff\xc0" + struct.pack(">HBHHB", 11, 8, 2, 3, 1) + b"\x01\x11\x00"
_SOS = b"\xff\xda" + struct.pack(">HB", 8, 1) + b"\x01\x00\x00\x3f\x00"


def _short_exif_jpeg(*, entries: int, length: int) -> bytes:
    """A 3x2 JFIF JPEG whose Exif block, after the frame header, says it is *length* bytes long.

    Its directory holds *entries* text entries, and no byte of the block is
    0xFF. Under 8 bytes, python-docx reads the block to the end of the file
    and parses it.
    """
    text = b"x" * 1023 + b"\x00"
    text_offset = 8 + 2 + 12 * entries + 4
    tags = [tag for tag in range(0x8001, 0x9000) if 0xFF not in tag.to_bytes(2, "big")][:entries]
    directory = (
        struct.pack(">H", entries)
        + b"".join(struct.pack(">HHII", tag, 2, len(text), text_offset) for tag in tags)
        + struct.pack(">I", 0)
    )
    tiff = b"MM\x00\x2a" + struct.pack(">I", 8) + directory + text
    app1 = b"\xff\xe1" + struct.pack(">H", length) + b"Exif\x00\x00" + tiff
    return b"\xff\xd8" + _JFIF_APP0 + _SOF0 + app1 + _SOS + b"\x00\xff\xd9"


def _padded_jpeg(*, markers: int, padding: bytes) -> bytes:
    """A 3x2 JFIF JPEG with *markers* in all before its scan, SOI first.

    *padding*, a TEM marker or an 0xFF 0x00 pair, fills them out after the
    frame header.
    """
    return b"\xff\xd8" + _JFIF_APP0 + _SOF0 + padding * (markers - 3) + _SOS + b"\x00\xff\xd9"


def _gif_signed(*, tag: bytes, jpeg: bytes) -> bytes:
    """*jpeg* behind a GIF signature, with *tag* at offset 6.

    python-docx checks for a JFIF or Exif tag at offset 6 before it checks
    for a GIF signature, so it reads this as a JPEG and walks every marker.
    """
    return b"GIF89a" + tag + b"\x00" + jpeg


def _chunky_png(*, chunks: int) -> bytes:
    """A 3x2 PNG of *chunks* chunks in all, padded with empty tEXt chunks."""
    png = _png(width=3, height=2)
    header, rest = png[:33], png[33:]  # the signature and IHDR, then IDAT and IEND
    empty_text = struct.pack(">I", 0) + b"tEXt" + struct.pack(">I", zlib.crc32(b"tEXt"))
    return header + empty_text * (chunks - 3) + rest


@pytest.mark.parametrize(
    ("data", "picture"),
    [
        (_exif_jpeg(entries=docx_writer._MAX_EXIF_ENTRIES), True),
        (_exif_jpeg(entries=docx_writer._MAX_EXIF_ENTRIES + 1), False),
        (_chunky_png(chunks=docx_writer._MAX_PNG_CHUNKS), True),
        (_chunky_png(chunks=docx_writer._MAX_PNG_CHUNKS + 1), False),
        (_short_exif_jpeg(entries=docx_writer._MAX_EXIF_ENTRIES + 1, length=2), False),
        (_short_exif_jpeg(entries=docx_writer._MAX_EXIF_ENTRIES + 1, length=7), False),
        (_padded_jpeg(markers=docx_writer._MAX_JPEG_MARKERS, padding=b"\xff\x01"), True),
        (_padded_jpeg(markers=docx_writer._MAX_JPEG_MARKERS + 1, padding=b"\xff\x01"), False),
        (_padded_jpeg(markers=docx_writer._MAX_JPEG_MARKERS + 1, padding=b"\xff\x00"), False),
        (
            _gif_signed(
                tag=b"JFIF",
                jpeg=_padded_jpeg(markers=docx_writer._MAX_JPEG_MARKERS + 1, padding=b"\xff\x01"),
            ),
            False,
        ),
        (
            _gif_signed(
                tag=b"Exif",
                jpeg=_short_exif_jpeg(entries=docx_writer._MAX_EXIF_ENTRIES + 1, length=2),
            ),
            False,
        ),
    ],
    ids=[
        "exif-at-the-cap", "exif-past-the-cap", "png-at-the-cap", "png-past-the-cap",
        "exif-shorter-than-its-header", "exif-a-byte-short-of-its-header",
        "markers-at-the-cap", "markers-past-the-cap", "ff00-pairs-past-the-cap",
        "gif-signed-jfif-past-the-marker-cap", "gif-signed-exif-shorter-than-its-header",
    ],
)
def test_a_header_python_docx_would_blow_up_on_becomes_the_placeholder(data, picture):
    mime = "image/png" if data.startswith(b"\x89PNG") else "image/jpeg"
    uri = f"data:{mime};base64," + base64.b64encode(data).decode()
    package = _write(f'<p>Chart: <img src="{uri}" alt="crafted"></p>')
    if picture:
        assert list(package.media().values()) == [data]
    else:
        assert package.media() == {}
        assert package.text == "Chart: [image: crafted]"


@pytest.mark.parametrize("case", _cases("image_sizes.yaml"), ids=_ids(_cases("image_sizes.yaml")))
def test_images_scale_down_to_fit(case):
    width_px, height_px = case["px"]
    image = f'<img src="{_png_data_uri(width=width_px, height=height_px)}" alt="generated">'
    if case["place"] == "cell":
        html = f"<table><tr><td>{image}</td><td>other</td></tr></table>"
    else:
        html = f"<p>{image}</p>"
    package = _write(html)
    [shape] = package.document.inline_shapes
    section = package.document.sections[0]
    native_width, native_height = width_px * 914_400 / 72, height_px * 914_400 / 72
    if case["fit"] == "native":
        assert (shape.width, shape.height) == (int(native_width), int(native_height))
        return
    if case["place"] == "cell":
        box = package.document.tables[0].cell(0, 0).width
    elif case["fit"] == "width":
        box = section.page_width - section.left_margin - section.right_margin
    else:
        box = section.page_height - section.top_margin - section.bottom_margin
    fitted, other, native_fitted, native_other = (
        (shape.width, shape.height, native_width, native_height)
        if case["fit"] == "width"
        else (shape.height, shape.width, native_height, native_width)
    )
    assert abs(fitted - box) <= 1
    assert abs(other - native_other * box / native_fitted) <= 1


def test_writer_imports_no_pillow():
    tree = ast.parse(Path(docx_writer.__file__).read_text(encoding="utf-8"))
    imported = {alias.name for node in ast.walk(tree) if isinstance(node, ast.Import) for alias in node.names}
    imported |= {node.module or "" for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)}
    assert not {name for name in imported if name == "PIL" or name.startswith("PIL.")}


def test_writer_takes_the_note_and_character_filter_from_the_inliner():
    # A second copy in the writer could drift from the HTML export's.
    tree = ast.parse(Path(docx_writer.__file__).read_text(encoding="utf-8"))
    literals = {node.value for node in ast.walk(tree) if isinstance(node, ast.Constant)}
    assert artifact_inline.TRUNCATION_NOTE not in literals
    assert artifact_inline.XML_UNSAFE.pattern not in literals
    assert docx_writer.TRUNCATION_NOTE is artifact_inline.TRUNCATION_NOTE
    assert docx_writer.XML_UNSAFE is artifact_inline.XML_UNSAFE


def test_writer_takes_the_image_signatures_and_placeholders_from_the_inliner():
    assert docx_writer.PNG_SIGNATURE is artifact_inline.PNG_SIGNATURE
    assert docx_writer.JPEG_SIGNATURE is artifact_inline.JPEG_SIGNATURE
    assert docx_writer.GIF_SIGNATURES is artifact_inline.GIF_SIGNATURES
    assert isinstance(docx_writer._SetAsideUris().placeholders, artifact_inline.UriPlaceholders)


def test_an_item_outside_any_list_is_a_first_level_bullet_and_deep_lists_stop_at_level_3():
    package = _write(
        "<li>loose</li>"
        "<ol><li>one<ul><li>two<ol><li>three<ul><li>four</li></ul></li></ol></li></ul></li></ol>"
    )
    styles = {p.text.strip(): p.style.name for p in package.document.paragraphs if p.text.strip()}

    assert styles == {
        "loose": "List Bullet",
        "one": "List Number",
        "two": "List Bullet 2",
        "three": "List Number 3",
        "four": "List Bullet 3",
    }


# Large inlined assets

LARGE_ASSETS = yaml.safe_load((FIXTURES / "large_assets.yaml").read_text(encoding="utf-8"))


def _large_png(*, index: int) -> bytes:
    return _padded_png(index=index, size=LARGE_ASSETS["png_bytes"])


def _png_uri(*, data: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(data).decode()


def _assets_html(*, case: dict, uri: Callable[[int], str]) -> str:
    assets = "".join(
        case["asset"].replace("{index}", str(index)).replace("{uri}", uri(index)) for index in range(case["count"])
    )
    return case["html"].replace("{assets}", assets)


def _assert_whole(*, package: Package, case: dict) -> None:
    """The heading, every picture in order, and the last paragraph survive, with no closing note."""
    texts = package.paragraph_texts()
    assert TRUNCATION_NOTE not in texts
    assert texts[-1] == "END"
    assert [p.text for p in package.document.paragraphs if p.style.name == "Heading 1"] == ["Report"]
    shapes = package.document.inline_shapes
    assert [shape._inline.docPr.get("descr") for shape in shapes] == [
        f"chart {index}" for index in range(case["pictures"])
    ]
    assert sorted(package.media().values()) == sorted(_large_png(index=index) for index in range(case["pictures"]))


@pytest.mark.parametrize("case", LARGE_ASSETS["writer"], ids=_ids(LARGE_ASSETS["writer"]))
def test_large_data_uris_keep_the_rest_of_the_document(case):
    package = _write(_assets_html(case=case, uri=lambda index: _png_uri(data=_large_png(index=index))))
    _assert_whole(package=package, case=case)


@pytest.fixture(scope="module")
def large_artifact(tmp_path_factory) -> Path:
    folder = tmp_path_factory.mktemp("large_artifact")
    for index in range(max(case["count"] for case in LARGE_ASSETS["inliner"])):
        (folder / f"big{index}.png").write_bytes(_large_png(index=index))
    return folder


@pytest.mark.parametrize("case", LARGE_ASSETS["inliner"], ids=_ids(LARGE_ASSETS["inliner"]))
def test_inliner_output_past_10_mb_converts_whole(case, large_artifact):
    html = _assets_html(case=case, uri=lambda index: f"big{index}.png")
    with pinned_dir(large_artifact) as folder:
        document = inline_artifact_html(
            html=html.encode(),
            folder=folder,
            base_parts=(),
            mode=InlineMode.WORD,
            limits=InlineLimits(max_total_bytes=case["budget"]),
        )
    with no_fetch() as recorder:
        data = write_docx(document=document, title="Report")
    assert recorder.calls == []
    _assert_whole(package=_open_package(data), case=case)


RESTORED_CASES = LARGE_ASSETS["restored"]["cases"]


@pytest.mark.parametrize("case", RESTORED_CASES, ids=_ids(RESTORED_CASES))
def test_long_data_uris_outside_images_read_back_unchanged(case):
    uri = _png_uri(data=_padded_png(index=0, size=LARGE_ASSETS["restored"]["png_bytes"]))
    package = _write(case["html"].replace("{uri}", uri))
    assert package.paragraph_texts() == [text.replace("{uri}", uri) for text in case["paragraphs"]]


# Hostile input


@pytest.mark.parametrize("case", _cases("hostile.yaml"), ids=_ids(_cases("hostile.yaml")))
def test_hostile_input_fetches_and_reads_nothing(case, tmp_path):
    sentinel = tmp_path / "sentinel.txt"
    sentinel.write_bytes(SENTINEL)
    with recording_server() as server:
        html = (
            case["html"]
            .replace("{server}", server.url)
            .replace("{sentinel_path}", str(sentinel))
            .replace("{sentinel_uri}", sentinel.as_uri())
        )
        with no_fetch() as recorder, _OPENS.recording() as opened:
            data = write_docx(document=_inlined(html), title="Hostile")
        hits = server.hits()
    assert recorder.calls == []
    assert _opened_outside_the_template(paths=opened) == []
    assert hits == []
    package = _open_package(data)
    assert not [name for name, body in package.parts.items() if SENTINEL in body]
    assert case["text"] in package.text
    assert package.media() == {}


@pytest.mark.parametrize("case", _cases("control_characters.yaml"), ids=_ids(_cases("control_characters.yaml")))
def test_xml_incompatible_characters_are_stripped(case):
    package = _write(case["html"])
    if "table_text" in case:
        assert package.document.tables[0].cell(0, 0).text == case["table_text"]
    else:
        # libxml2 before 2.12 cuts an attribute value at an invalid character
        # reference instead of reading past it.
        before_2_12 = etree.LIBXML_VERSION < (2, 12, 0)
        expected = case.get("text_before_libxml2_2_12", case["text"]) if before_2_12 else case["text"]
        assert [p.text for p in package.document.paragraphs] == [expected]


# Tables


def test_tables_nest_only_as_deep_as_word_reads_and_the_deepest_text_stays():
    nested = 40
    package = _write("<table><tr><td>" * nested + "deep text" + "</td></tr></table>" * nested)
    body = etree.fromstring(package.parts["word/document.xml"], _SAFE_XML)
    tables = [element for element in body.iter() if element.tag.endswith("}tbl")]
    depth = max(sum(1 for a in table.iterancestors() if a.tag.endswith("}tbl")) for table in tables) + 1
    assert depth == docx_writer.MAX_TABLE_NESTING
    assert "deep text" in package.text


@pytest.mark.parametrize("case", _cases("tables.yaml"), ids=_ids(_cases("tables.yaml")))
def test_table_layouts(case):
    package = _write(case["html"].replace("{digits}", _HUGE_NUMBER))
    [table] = package.document.tables
    assert table.style.name == "Table Grid"
    assert (len(table.rows), len(table.columns)) == (case["rows"], case["columns"])
    for start, end in case["merged"]:
        assert table.cell(*start)._tc is table.cell(*end)._tc
    for text in case["present"]:
        assert text in package.text
    for text in case["absent"]:
        assert text not in package.text


@pytest.mark.parametrize("case", _cases("table_content.yaml"), ids=_ids(_cases("table_content.yaml")))
def test_content_outside_table_cells_is_kept(case):
    package = _write(case["html"])
    assert package.paragraph_texts() == case["paragraphs"]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in package.document.tables]
    assert tables == case["tables"]


def test_tables_clamp_to_63_columns_and_say_so():
    cells = "".join(f"<td>c{index}</td>" for index in range(1, MAX_TABLE_COLUMNS + 2))
    package = _write(f"<table><tr>{cells}</tr></table>")
    [table] = package.document.tables
    assert len(table.columns) == MAX_TABLE_COLUMNS
    assert table.cell(0, MAX_TABLE_COLUMNS - 1).text == f"c{MAX_TABLE_COLUMNS}"
    assert f"c{MAX_TABLE_COLUMNS + 1}" not in package.text
    assert package.document.paragraphs[-1].text == TABLE_CELLS_NOTE


def test_a_cell_pushed_past_63_columns_by_spans_is_left_out_and_said_so():
    spanning = f'<td colspan="{MAX_TABLE_COLUMNS}">wide</td><td>pushed out</td>'
    package = _write(f"<table><tr>{spanning}</tr></table>")
    assert "pushed out" not in package.text
    assert package.document.paragraphs[-1].text == TABLE_CELLS_NOTE


def test_rows_past_the_cell_budget_are_left_out_and_said_so():
    row = "<tr>" + "<td>cell</td>" * MAX_TABLE_COLUMNS + "</tr>"
    rows_that_fit = MAX_TABLE_CELLS // MAX_TABLE_COLUMNS
    package = _write(f"<table>{row * (rows_that_fit + 1)}</table><p>after</p>")
    [table] = package.document.tables
    assert len(table.rows) == rows_that_fit
    assert [paragraph.text for paragraph in package.document.paragraphs[-2:]] == [
        "after",
        TABLE_CELLS_NOTE,
    ]


@pytest.mark.parametrize(
    "late", ["<tr><td>late</td></tr>", "<td>late</td>"], ids=["row", "loose_cell"]
)
def test_a_table_after_the_cell_budget_is_spent_is_left_out_and_said_so(late):
    full_row = "<tr>" + "<td></td>" * MAX_TABLE_COLUMNS + "</tr>"
    rows_that_fit = MAX_TABLE_CELLS // MAX_TABLE_COLUMNS
    rest = "<tr><td></td></tr>" * (MAX_TABLE_CELLS - rows_that_fit * MAX_TABLE_COLUMNS)
    package = _write(
        f"<table>{full_row * rows_that_fit}</table><table>{rest}</table><table>{late}</table>"
    )
    assert len(package.document.tables) == 2
    assert "late" not in package.text
    assert package.document.paragraphs[-1].text == TABLE_CELLS_NOTE


def test_a_table_within_the_limits_says_nothing():
    package = _write("<table><tr><td>a</td><td>b</td></tr></table><p>after</p>")
    assert TABLE_CELLS_NOTE not in package.text
    assert package.document.paragraphs[-1].text == "after"


def test_table_cells_stay_within_the_document_budget():
    full_row = "<tr>" + "<td></td>" * MAX_TABLE_COLUMNS + "</tr>"
    rows_that_fit = MAX_TABLE_CELLS // MAX_TABLE_COLUMNS
    spanning_row = f'<tr><td colspan="{MAX_TABLE_COLUMNS}" rowspan="65534">span</td></tr>'
    # The first table spends all but MAX_TABLE_CELLS % MAX_TABLE_COLUMNS
    # cells, too few for one row of the second.
    second_row = "<tr>" + "<td>second table</td>" * MAX_TABLE_COLUMNS + "</tr>"
    html = (
        f"<table>{spanning_row}{full_row * (rows_that_fit + 10)}</table>"
        f"<table>{second_row}</table>"
        "<p>after the tables</p>"
    )
    package = _write(html)
    [table] = package.document.tables
    assert (len(table.rows), len(table.columns)) == (rows_that_fit, MAX_TABLE_COLUMNS)
    assert table.cell(0, 0)._tc is table.cell(rows_that_fit - 1, MAX_TABLE_COLUMNS - 1)._tc
    assert len(list(package.document.element.body.iter(qn("w:tc")))) <= MAX_TABLE_CELLS
    assert "second table" not in package.text
    assert [paragraph.text for paragraph in package.document.paragraphs[-2:]] == [
        "after the tables",
        TABLE_CELLS_NOTE,
    ]


def test_nested_table_keeps_the_text_order():
    package = _write("<table><tr><td>outer<table><tr><td>inner</td></tr></table>after</td></tr></table>")
    [table] = package.document.tables
    cell = table.cell(0, 0)
    assert [nested.cell(0, 0).text for nested in cell.tables] == ["inner"]
    assert package.text.split("\n") == ["outer", "inner", "after"]


# Walk limits


def test_limits_hold_their_specified_values():
    limits = yaml.safe_load((FIXTURES / "limits.yaml").read_text(encoding="utf-8"))
    assert {
        "max_depth": MAX_DEPTH,
        "max_nodes": MAX_NODES,
        "deadline_seconds": DEADLINE_SECONDS,
        "max_table_columns": MAX_TABLE_COLUMNS,
        "max_table_cells": MAX_TABLE_CELLS,
    } == limits


@pytest.mark.parametrize("separator", ["\n", "\t"], ids=["line_breaks", "tabs"])
def test_line_breaks_and_tabs_in_pre_spend_the_node_budget(separator):
    package = _write("<pre>head" + separator * (MAX_NODES + 10) + "tail</pre>")
    body = package.document.element.body
    assert sum(1 for _ in body.iter(qn("w:br"), qn("w:tab"))) <= MAX_NODES
    assert package.document.paragraphs[-1].text == TRUNCATION_NOTE
    # python-docx reads a paragraph's text slowly past many line breaks, so
    # this reads the text elements directly.
    assert "tail" not in "".join(t.text for t in body.iter(qn("w:t")))


@pytest.mark.parametrize("case", _cases("wide_elements.yaml"), ids=_ids(_cases("wide_elements.yaml")))
def test_children_are_queued_one_at_a_time(case, monkeypatch):
    queued = 0

    def counting(cls: type) -> object:
        def make(**fields: object) -> object:
            nonlocal queued
            queued += 1
            return cls(**fields)

        return make

    monkeypatch.setattr(docx_writer, "MAX_NODES", case["budget"])
    monkeypatch.setattr(docx_writer, "_Visit", counting(docx_writer._Visit))
    monkeypatch.setattr(docx_writer, "_TableItem", counting(docx_writer._TableItem))
    package = _write(case["open"] + case["child"] * case["repeat"] + case["close"])
    assert package.document.paragraphs[-1].text == TRUNCATION_NOTE
    assert queued <= 4 * case["budget"]


def test_node_budget_ends_with_the_closing_paragraph():
    package = _write("<p>first</p>" + "<i></i>" * MAX_NODES + "<p>past the budget</p>")
    paragraphs = package.document.paragraphs
    assert paragraphs[0].text == "first"
    assert paragraphs[-1].text == TRUNCATION_NOTE
    assert "past the budget" not in package.text


def test_deadline_ends_with_the_closing_paragraph(monkeypatch):
    readings: list[float] = []

    def monotonic() -> float:
        # The deadline is set from the first reading; the sixth is past it.
        readings.append(0.0 if len(readings) < 5 else DEADLINE_SECONDS + 1)
        return readings[-1]

    monkeypatch.setattr(docx_writer, "time", types.SimpleNamespace(monotonic=monotonic))
    package = _write("".join(f"<p>paragraph {index}</p>" for index in range(50)))
    texts = [p.text for p in package.document.paragraphs]
    assert texts[0] == "paragraph 0"
    assert texts[-1] == TRUNCATION_NOTE
    assert "paragraph 49" not in texts
    assert len(readings) == 6


@pytest.mark.parametrize("case", _cases("deep_nesting.yaml"), ids=_ids(_cases("deep_nesting.yaml")))
def test_deep_nesting_raises_no_recursion_error(case):
    repeat = case["repeat"]
    package = _write(case["open"] * repeat + "deep text" + case["close"] * repeat + "<p>after</p>")
    assert package.document.paragraphs[-1].text == TRUNCATION_NOTE


def test_past_the_depth_limit_text_still_flows(monkeypatch):
    monkeypatch.setattr(docx_writer, "MAX_DEPTH", 6)
    nested = "<div>" * 10 + "<h1>deep heading</h1><table><tr><td>deep cell</td></tr></table>" + "</div>" * 10
    package = _write(nested + "<p>after</p>")
    assert package.document.tables == []
    assert not [p for p in package.document.paragraphs if p.style.name.startswith("Heading")]
    assert "deep heading" in package.text
    assert "deep cell" in package.text
